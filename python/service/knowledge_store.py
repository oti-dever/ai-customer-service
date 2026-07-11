from __future__ import annotations

import hashlib
import base64
import json
import logging
import math
import mimetypes
import os
import re
import shutil
import sqlite3
import struct
import sys
import time
import uuid
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from rpa.db.connection import PROJECT_ROOT, open_db

PROJECT_PYDEPS_DIR = PROJECT_ROOT / ".pydeps"
if PROJECT_PYDEPS_DIR.is_dir():
    project_pydeps = str(PROJECT_PYDEPS_DIR)
    if project_pydeps not in sys.path:
        sys.path.append(project_pydeps)

SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf", ".xlsx", ".xlsm"}
SUPPORTED_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif"}
MAX_EXCEL_ROWS_PER_SHEET = 2000
MAX_EXCEL_COLUMNS_PER_SHEET = 80
MAX_EXCEL_CELL_CHARS = 500
DEFAULT_KNOWLEDGE_DB_PATH = PROJECT_ROOT / "database" / "knowledge_store.db"
DEFAULT_KNOWLEDGE_FILES_DIR = PROJECT_ROOT / "database" / "knowledge_files"
LOCAL_EMBEDDING_MODEL = "local-hash-ngram-v1"
LOCAL_EMBEDDING_DIM = 384
BGE_SMALL_ZH_MODEL = "BAAI/bge-small-zh-v1.5"
BGE_SMALL_ZH_LOCAL_DIR = PROJECT_ROOT / "database" / "models" / "bge-small-zh-v1.5"
BGE_SMALL_ZH_DIM = 512
BGE_VL_MODEL = "BAAI/bge-vl-base"
BGE_VL_LOCAL_DIR = PROJECT_ROOT / "database" / "models" / "bge-vl-base"
CHINESE_CLIP_MODEL = "OFA-Sys/chinese-clip-vit-base-patch16"
CHINESE_CLIP_LOCAL_DIR = PROJECT_ROOT / "database" / "models" / "chinese-clip-vit-base-patch16"
BGE_QUERY_INSTRUCTION = "为这个句子生成表示以用于检索相关文章："
VECTOR_SCORE_THRESHOLD = 0.05
DEFAULT_VECTOR_WEIGHT = 0.6
DEFAULT_KEYWORD_WEIGHT = 0.4
DEFAULT_IMAGE_VECTOR_WEIGHT = 0.35
DEFAULT_IMAGE_TEXT_WEIGHT = 0.25
DEFAULT_IMAGE_RULE_WEIGHT = 0.4
ADJACENT_CHUNK_MERGE_RADIUS = 1
ADJACENT_CHUNK_MERGE_MAX_CHARS = 1600
ARK_IMAGE_ANALYSIS_DEFAULT_BASE_URL = "https://ark.cn-beijing.volces.com/api/v3"
ARK_IMAGE_ANALYSIS_DEFAULT_MODEL = ""
IMAGE_ANALYSIS_PROMPT = """你是电商客服图片素材入库分析助手。请分析这张店铺客服可能使用的图片素材，只输出合法 JSON，不要 Markdown，不要解释。
字段固定如下：
{
  "asset_type": "商品图/实拍图/规格图/尺寸图/安装说明/售后说明/物流说明/活动图/纯文本图/其他",
  "summary": "一句话说明图片主要内容，面向内部素材管理，不直接发给客户",
  "ocr_text": "图片中能读到的文字，没有则为空字符串",
  "tags": ["3到8个检索标签"],
  "scenarios": ["适合使用的客服场景，如外观咨询、尺寸咨询、安装咨询、售后咨询"],
  "risk_tags": ["风险标签，如含价格、含活动、含过期信息、含平台水印、含二维码、含隐私信息、不建议自动发送；无风险则空数组"],
  "suggested_reply": "客服发送该图时可搭配的一句自然话术，没有则为空字符串"
}
要求：不确定的具体商品型号、价格、活动不要编造；发现价格、活动、二维码、水印、隐私信息时必须写入 risk_tags。"""


class EmbeddingProviderError(RuntimeError):
    pass


class ImageAssetAnalyzer:
    def analyze(self, path: Path) -> dict[str, Any]:
        raise NotImplementedError


class LocalImageAssetAnalyzer(ImageAssetAnalyzer):
    def analyze(self, path: Path) -> dict[str, Any]:
        width, height = _read_image_size(path)
        suffix = path.suffix.lower().lstrip(".")
        title = _humanize_asset_name(path.stem)
        tags = [tag for tag in _query_terms(title) if len(tag) >= 2][:8]
        if suffix:
            tags.append(suffix)
        summary_parts = [f"本地图片素材：{title}"]
        if width > 0 and height > 0:
            summary_parts.append(f"图片尺寸约 {width}x{height}")
        summary = "，".join(summary_parts)
        return {
            "asset_type": "other",
            "summary": summary,
            "ocr_text": "",
            "tags": sorted(set(tags)),
            "scenarios": [],
            "risk_tags": [],
            "suggested_reply": "",
            "analysis_model": "local-image-fallback-v1",
            "width": width,
            "height": height,
            "raw_analysis": {"source": "local_fallback", "filename": path.name},
        }


class ArkDoubaoImageAssetAnalyzer(ImageAssetAnalyzer):
    def __init__(
        self,
        api_key: str,
        model: str,
        base_url: str = ARK_IMAGE_ANALYSIS_DEFAULT_BASE_URL,
        timeout: float = 45.0,
        transport: Callable[[Request, float], bytes] | None = None,
    ) -> None:
        self.api_key = api_key.strip()
        self.model = model.strip()
        self.base_url = base_url.strip().rstrip("/") or ARK_IMAGE_ANALYSIS_DEFAULT_BASE_URL
        self.timeout = timeout
        self._transport = transport or self._default_transport
        if not self.api_key:
            raise ValueError("missing image analysis api key")
        if not self.model:
            raise ValueError("missing image analysis model")

    def analyze(self, path: Path) -> dict[str, Any]:
        data_url = _image_file_to_data_url(path)
        payload = {
            "model": self.model,
            "stream": False,
            "messages": [
                {"role": "system", "content": IMAGE_ANALYSIS_PROMPT},
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": "请分析这张客服图片素材，并按指定 JSON 字段返回。"},
                        {"type": "image_url", "image_url": {"url": data_url}},
                    ],
                },
            ],
            "thinking": {"type": "disabled"},
            "temperature": 0.1,
        }
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        request = Request(
            _chat_completions_url(self.base_url),
            data=body,
            method="POST",
            headers={
                "Content-Type": "application/json",
                "Authorization": f"Bearer {self.api_key}",
            },
        )
        try:
            response_body = self._transport(request, self.timeout)
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")[:800]
            raise RuntimeError(f"ark image analysis http {exc.code}: {detail}") from exc
        except URLError as exc:
            raise RuntimeError(f"ark image analysis network error: {exc.reason}") from exc

        content = _extract_chat_completion_content(response_body)
        analysis = _extract_json_object_from_text(content)
        normalized = _normalize_image_analysis(analysis)
        normalized["analysis_model"] = f"ark:{self.model}"
        normalized["raw_analysis"] = analysis
        return normalized

    @staticmethod
    def _default_transport(request: Request, timeout: float) -> bytes:
        with urlopen(request, timeout=timeout) as response:
            return response.read()


class EmbeddingProvider:
    model_name: str
    dim: int

    def encode(self, text: str, is_query: bool = False) -> list[float]:
        raise NotImplementedError


class ImageEmbeddingProvider:
    model_name: str
    dim: int

    def encode_text(self, text: str) -> list[float]:
        raise NotImplementedError

    def encode_image(self, path: Path) -> list[float]:
        raise NotImplementedError


class LocalHashEmbeddingProvider(EmbeddingProvider):
    model_name = LOCAL_EMBEDDING_MODEL
    dim = LOCAL_EMBEDDING_DIM

    def encode(self, text: str, is_query: bool = False) -> list[float]:
        del is_query
        return _local_embedding(text, self.dim)


class SentenceTransformerEmbeddingProvider(EmbeddingProvider):
    def __init__(self, model_name: str, canonical_model_name: str | None = None) -> None:
        try:
            from sentence_transformers import SentenceTransformer
        except ImportError as exc:  # pragma: no cover - depends on optional runtime packages
            raise EmbeddingProviderError("sentence-transformers is not installed") from exc

        device = os.environ.get("AI_CUSTOMER_SERVICE_EMBEDDING_DEVICE", "cpu").strip() or "cpu"
        cache_dir = os.environ.get("AI_CUSTOMER_SERVICE_EMBEDDING_CACHE_DIR", "").strip() or None
        allow_download = os.environ.get("AI_CUSTOMER_SERVICE_EMBEDDING_ALLOW_DOWNLOAD", "").strip() == "1"
        try:
            self._model = SentenceTransformer(
                model_name,
                device=device,
                cache_folder=cache_dir,
                local_files_only=not allow_download,
            )
        except Exception as exc:  # pragma: no cover - depends on local model/network state
            raise EmbeddingProviderError(f"failed to load embedding model {model_name}: {exc}") from exc
        self.model_name = canonical_model_name or model_name
        if hasattr(self._model, "get_embedding_dimension"):
            self.dim = int(self._model.get_embedding_dimension() or 0)
        else:
            self.dim = int(self._model.get_sentence_embedding_dimension() or 0)

    def encode(self, text: str, is_query: bool = False) -> list[float]:
        text = text.strip()
        if is_query and self.model_name == BGE_SMALL_ZH_MODEL:
            text = f"{BGE_QUERY_INSTRUCTION}{text}"
        if not text:
            return [0.0] * self.dim
        vector = self._model.encode(
            [text],
            normalize_embeddings=True,
            convert_to_numpy=True,
            show_progress_bar=False,
        )[0]
        return [float(value) for value in vector.tolist()]


class SentenceTransformerImageEmbeddingProvider(ImageEmbeddingProvider):
    def __init__(self, model_name: str, canonical_model_name: str | None = None) -> None:
        try:
            from sentence_transformers import SentenceTransformer
        except ImportError as exc:  # pragma: no cover - depends on optional runtime packages
            raise EmbeddingProviderError("sentence-transformers is not installed") from exc

        device = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_EMBEDDING_DEVICE", "cpu").strip() or "cpu"
        cache_dir = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_EMBEDDING_CACHE_DIR", "").strip() or None
        allow_download = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_EMBEDDING_ALLOW_DOWNLOAD", "").strip() == "1"
        model_path = Path(model_name)
        if model_path.exists():
            model_dir = str(model_path.resolve())
            if model_dir not in sys.path:
                sys.path.insert(0, model_dir)
        try:
            self._model = SentenceTransformer(
                model_name,
                device=device,
                cache_folder=cache_dir,
                local_files_only=not allow_download,
                trust_remote_code=True,
            )
        except Exception as exc:  # pragma: no cover - depends on local model/network state
            raise EmbeddingProviderError(f"failed to load image embedding model {model_name}: {exc}") from exc
        self.model_name = canonical_model_name or model_name
        if hasattr(self._model, "get_embedding_dimension"):
            self.dim = int(self._model.get_embedding_dimension() or 0)
        else:
            self.dim = 0

    def _encode_value(self, value: Any) -> list[float]:
        vector = self._model.encode(
            [value],
            normalize_embeddings=True,
            convert_to_numpy=True,
            show_progress_bar=False,
        )[0]
        values = [float(item) for item in vector.tolist()]
        if not self.dim:
            self.dim = len(values)
        return values

    def encode_text(self, text: str) -> list[float]:
        text = text.strip()
        if not text:
            return [0.0] * max(0, self.dim)
        return self._encode_value(text)

    def encode_image(self, path: Path) -> list[float]:
        return self._encode_value(str(path))


class TransformersImageEmbeddingProvider(ImageEmbeddingProvider):
    def __init__(self, model_name: str, canonical_model_name: str | None = None) -> None:
        try:
            import torch
            from PIL import Image
            from transformers import AutoModel, AutoProcessor, CLIPModel, CLIPProcessor
        except ImportError as exc:  # pragma: no cover - depends on optional runtime packages
            raise EmbeddingProviderError("transformers/torch/pillow is not installed") from exc

        self._torch = torch
        self._image_cls = Image
        self.model_name = canonical_model_name or model_name
        device = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_EMBEDDING_DEVICE", "cpu").strip() or "cpu"
        cache_dir = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_EMBEDDING_CACHE_DIR", "").strip() or None
        allow_download = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_EMBEDDING_ALLOW_DOWNLOAD", "").strip() == "1"
        try:
            try:
                self._processor = AutoProcessor.from_pretrained(
                    model_name,
                    cache_dir=cache_dir,
                    local_files_only=not allow_download,
                )
            except Exception:
                self._processor = CLIPProcessor.from_pretrained(
                    model_name,
                    cache_dir=cache_dir,
                    local_files_only=not allow_download,
                )
            try:
                self._model = AutoModel.from_pretrained(
                    model_name,
                    cache_dir=cache_dir,
                    local_files_only=not allow_download,
                )
            except Exception:
                self._model = CLIPModel.from_pretrained(
                    model_name,
                    cache_dir=cache_dir,
                    local_files_only=not allow_download,
                )
        except Exception as exc:  # pragma: no cover - depends on local model/network state
            raise EmbeddingProviderError(f"failed to load image embedding model {model_name}: {exc}") from exc
        self._model.to(device)
        self._model.eval()
        self._device = device
        self.dim = self._infer_dim()

    def _infer_dim(self) -> int:
        for attr in ("projection_dim", "hidden_size"):
            value = getattr(getattr(self._model, "config", object()), attr, 0)
            if int(value or 0) > 0:
                return int(value)
        return 0

    def _normalize_vector(self, tensor: Any) -> list[float]:
        vector = tensor.detach().cpu().float()
        if len(vector.shape) > 1:
            vector = vector[0]
        norm = self._torch.linalg.vector_norm(vector)
        if float(norm) > 0:
            vector = vector / norm
        values = [float(value) for value in vector.tolist()]
        if not self.dim:
            self.dim = len(values)
        return values

    def encode_text(self, text: str) -> list[float]:
        text = text.strip()
        if not text:
            return [0.0] * max(0, self.dim)
        inputs = self._processor(text=[text], return_tensors="pt", padding=True, truncation=True)
        inputs = {key: value.to(self._device) for key, value in inputs.items()}
        with self._torch.no_grad():
            if hasattr(self._model, "get_text_features"):
                output = self._model.get_text_features(**inputs)
            else:
                output = self._model(**inputs)
                output = getattr(output, "text_embeds", None) or getattr(output, "pooler_output", None)
            if output is None:
                raise EmbeddingProviderError(f"model {self.model_name} does not expose text embeddings")
            return self._normalize_vector(output)

    def encode_image(self, path: Path) -> list[float]:
        with self._image_cls.open(path) as image:
            rgb_image = image.convert("RGB")
            inputs = self._processor(images=rgb_image, return_tensors="pt")
        inputs = {key: value.to(self._device) for key, value in inputs.items()}
        with self._torch.no_grad():
            if hasattr(self._model, "get_image_features"):
                output = self._model.get_image_features(**inputs)
            else:
                output = self._model(**inputs)
                output = getattr(output, "image_embeds", None) or getattr(output, "pooler_output", None)
            if output is None:
                raise EmbeddingProviderError(f"model {self.model_name} does not expose image embeddings")
            return self._normalize_vector(output)


_embedding_providers: dict[str, EmbeddingProvider] = {}
_image_embedding_providers: dict[str, ImageEmbeddingProvider] = {}


def resolved_knowledge_db_path() -> Path:
    raw = os.environ.get("AI_CUSTOMER_SERVICE_KNOWLEDGE_DB", "").strip()
    if raw:
        return Path(raw).expanduser()
    return DEFAULT_KNOWLEDGE_DB_PATH


def resolved_knowledge_files_dir() -> Path:
    raw = os.environ.get("AI_CUSTOMER_SERVICE_KNOWLEDGE_FILES", "").strip()
    if raw:
        return Path(raw).expanduser()
    return DEFAULT_KNOWLEDGE_FILES_DIR


def _open_knowledge_db(db_path: Path) -> sqlite3.Connection:
    conn = open_db(db_path)
    conn.row_factory = sqlite3.Row
    return conn


def _now_text() -> str:
    return time.strftime("%Y-%m-%d %H:%M:%S")


def _new_id(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex}"


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _chat_completions_url(base_url: str) -> str:
    value = base_url.strip().rstrip("/")
    suffix = "/chat/completions"
    if value.endswith(suffix):
        return value
    return f"{value}{suffix}"


def _image_file_to_data_url(path: Path) -> str:
    mime_type = mimetypes.guess_type(str(path))[0] or "application/octet-stream"
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def _extract_chat_completion_content(body: bytes) -> str:
    try:
        payload = json.loads(body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise RuntimeError("invalid chat completion json response") from exc
    if not isinstance(payload, dict):
        raise RuntimeError("chat completion response is not an object")
    choices = payload.get("choices")
    if not isinstance(choices, list) or not choices:
        raise RuntimeError("chat completion response has no choices")
    first = choices[0]
    if not isinstance(first, dict):
        raise RuntimeError("chat completion choice is not an object")
    message = first.get("message")
    if isinstance(message, dict):
        content = message.get("content")
        if isinstance(content, str) and content.strip():
            return content.strip()
    delta = first.get("delta")
    if isinstance(delta, dict):
        content = delta.get("content")
        if isinstance(content, str) and content.strip():
            return content.strip()
    raise RuntimeError("chat completion response has empty content")


def _extract_json_object_from_text(text: str) -> dict[str, Any]:
    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
        cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        value = json.loads(cleaned)
    except json.JSONDecodeError:
        start = cleaned.find("{")
        end = cleaned.rfind("}")
        if start < 0 or end <= start:
            raise RuntimeError("image analysis response is not json")
        value = json.loads(cleaned[start : end + 1])
    if not isinstance(value, dict):
        raise RuntimeError("image analysis json is not an object")
    return value


def _normalize_image_analysis(value: dict[str, Any]) -> dict[str, Any]:
    asset_type = str(value.get("asset_type") or "other").strip() or "other"
    summary = str(value.get("summary") or "").strip()
    ocr_text = str(value.get("ocr_text") or "").strip()
    tags = _string_list(value.get("tags"), limit=16)
    scenarios = _string_list(value.get("scenarios"), limit=16)
    risk_tags = _string_list(value.get("risk_tags"), limit=16)
    suggested_reply = str(value.get("suggested_reply") or "").strip()
    return {
        "asset_type": asset_type,
        "summary": summary,
        "ocr_text": ocr_text,
        "tags": tags,
        "scenarios": scenarios,
        "risk_tags": risk_tags,
        "suggested_reply": suggested_reply,
        "analysis_model": str(value.get("analysis_model") or "").strip(),
        "raw_analysis": value,
    }


def _json_dumps(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _json_loads_object(text: str) -> dict[str, Any]:
    try:
        value = json.loads(text or "{}")
    except json.JSONDecodeError:
        return {}
    return value if isinstance(value, dict) else {}


def _json_loads_list(text: str) -> list[str]:
    try:
        value = json.loads(text or "[]")
    except json.JSONDecodeError:
        return []
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if str(item).strip()]


def _string_list(value: Any, limit: int = 16) -> list[str]:
    if isinstance(value, str):
        items = re.split(r"[,，、\s]+", value)
    elif isinstance(value, (list, tuple, set)):
        items = [str(item) for item in value]
    else:
        items = []
    result: list[str] = []
    seen: set[str] = set()
    for item in items:
        cleaned = str(item).strip()
        if not cleaned or cleaned in seen:
            continue
        seen.add(cleaned)
        result.append(cleaned)
        if len(result) >= limit:
            break
    return result


def _humanize_asset_name(name: str) -> str:
    text = re.sub(r"[_\-]+", " ", name).strip()
    return text or name


def _read_image_size(path: Path) -> tuple[int, int]:
    data = path.read_bytes()[:256]
    if data.startswith(b"\x89PNG\r\n\x1a\n") and len(data) >= 24:
        return int.from_bytes(data[16:20], "big"), int.from_bytes(data[20:24], "big")
    if data.startswith(b"GIF87a") or data.startswith(b"GIF89a"):
        if len(data) >= 10:
            return int.from_bytes(data[6:8], "little"), int.from_bytes(data[8:10], "little")
    if data.startswith(b"BM") and len(data) >= 26:
        return int.from_bytes(data[18:22], "little", signed=True), int.from_bytes(data[22:26], "little", signed=True)
    if data.startswith(b"RIFF") and data[8:12] == b"WEBP":
        return _read_webp_size(path)
    if data.startswith(b"\xff\xd8"):
        return _read_jpeg_size(path)
    return 0, 0


def _read_jpeg_size(path: Path) -> tuple[int, int]:
    data = path.read_bytes()
    index = 2
    while index + 9 < len(data):
        if data[index] != 0xFF:
            index += 1
            continue
        marker = data[index + 1]
        index += 2
        if marker in {0xD8, 0xD9}:
            continue
        if index + 2 > len(data):
            break
        length = int.from_bytes(data[index : index + 2], "big")
        if length < 2 or index + length > len(data):
            break
        if marker in {0xC0, 0xC1, 0xC2, 0xC3, 0xC5, 0xC6, 0xC7, 0xC9, 0xCA, 0xCB, 0xCD, 0xCE, 0xCF}:
            if index + 7 <= len(data):
                height = int.from_bytes(data[index + 3 : index + 5], "big")
                width = int.from_bytes(data[index + 5 : index + 7], "big")
                return width, height
        index += length
    return 0, 0


def _read_webp_size(path: Path) -> tuple[int, int]:
    data = path.read_bytes()
    if len(data) < 30 or data[:4] != b"RIFF" or data[8:12] != b"WEBP":
        return 0, 0
    chunk = data[12:16]
    if chunk == b"VP8X" and len(data) >= 30:
        width = int.from_bytes(data[24:27], "little") + 1
        height = int.from_bytes(data[27:30], "little") + 1
        return width, height
    if chunk == b"VP8 " and len(data) >= 30:
        width = int.from_bytes(data[26:28], "little") & 0x3FFF
        height = int.from_bytes(data[28:30], "little") & 0x3FFF
        return width, height
    if chunk == b"VP8L" and len(data) >= 25:
        bits = int.from_bytes(data[21:25], "little")
        width = (bits & 0x3FFF) + 1
        height = ((bits >> 14) & 0x3FFF) + 1
        return width, height
    return 0, 0


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\t\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - depends on runtime packages
        raise RuntimeError("python-docx is required to extract docx files") from exc

    doc = Document(str(path))
    pieces: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            pieces.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                pieces.append(row_text)
    return "\n".join(pieces)


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on runtime packages
        raise RuntimeError("pypdf is required to extract pdf files") from exc

    reader = PdfReader(str(path))
    pieces = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pieces.append(text)
    return "\n".join(pieces)


def _excel_cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if hasattr(value, "isoformat"):
        try:
            return str(value.isoformat())
        except Exception:
            pass
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) > MAX_EXCEL_CELL_CHARS:
        return text[:MAX_EXCEL_CELL_CHARS] + "..."
    return text


def _trim_excel_row(cells: list[str]) -> list[str]:
    end = len(cells)
    while end > 0 and not cells[end - 1]:
        end -= 1
    return cells[:end]


def _find_excel_header_row(rows: list[tuple[int, list[str]]]) -> int:
    candidates = rows[: min(20, len(rows))]
    if not candidates:
        return 0
    best_index = 0
    best_score = -1
    for index, (_row_number, cells) in enumerate(candidates):
        non_empty = sum(1 for cell in cells if cell)
        text_score = sum(1 for cell in cells if cell and not re.fullmatch(r"[-+]?\d+(\.\d+)?", cell))
        score = non_empty * 10 + text_score
        if score > best_score:
            best_score = score
            best_index = index
    return best_index


def _extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on runtime packages
        raise RuntimeError("openpyxl is required to extract Excel files") from exc

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    pieces: list[str] = []
    try:
        for sheet in workbook.worksheets:
            max_row = min(sheet.max_row or 0, MAX_EXCEL_ROWS_PER_SHEET)
            max_col = min(sheet.max_column or 0, MAX_EXCEL_COLUMNS_PER_SHEET)
            rows: list[tuple[int, list[str]]] = []
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True),
                start=1,
            ):
                cells = _trim_excel_row([_excel_cell_to_text(value) for value in row])
                if any(cells):
                    rows.append((row_number, cells))
            if not rows:
                continue

            pieces.append(f"【工作表：{sheet.title}】")
            if sheet.sheet_state != "visible":
                pieces.append(f"工作表状态：{sheet.sheet_state}")

            header_index = _find_excel_header_row(rows)
            for row_number, cells in rows[:header_index]:
                row_text = " | ".join(cell for cell in cells if cell)
                if row_text:
                    pieces.append(f"说明行{row_number}：{row_text}")

            header_row_number, header_cells = rows[header_index]
            del header_row_number
            headers: list[str] = []
            for column_index, header in enumerate(header_cells, start=1):
                headers.append(header or f"第{column_index}列")
            pieces.append("表头：" + " | ".join(headers))
            pieces.append("")

            for row_number, cells in rows[header_index + 1 :]:
                row_pairs: list[str] = []
                for column_index, value in enumerate(cells):
                    if not value:
                        continue
                    header = headers[column_index] if column_index < len(headers) else f"第{column_index + 1}列"
                    row_pairs.append(f"{header}={value}")
                if row_pairs:
                    pieces.append(f"第{row_number}行：")
                    pieces.extend(row_pairs)
                    pieces.append("")

            if (sheet.max_row or 0) > MAX_EXCEL_ROWS_PER_SHEET:
                pieces.append(f"提示：工作表超过 {MAX_EXCEL_ROWS_PER_SHEET} 行，导入文本已截断。")
            if (sheet.max_column or 0) > MAX_EXCEL_COLUMNS_PER_SHEET:
                pieces.append(f"提示：工作表超过 {MAX_EXCEL_COLUMNS_PER_SHEET} 列，导入文本已截断。")
            pieces.append("")
    finally:
        workbook.close()
    return "\n".join(pieces)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return _clean_text(_decode_text(path.read_bytes()))
    if suffix == ".docx":
        return _clean_text(_extract_docx(path))
    if suffix == ".pdf":
        return _clean_text(_extract_pdf(path))
    if suffix in {".xlsx", ".xlsm"}:
        return _clean_text(_extract_xlsx(path))
    raise ValueError(f"unsupported knowledge document type: {suffix}")


def _heading_title(line: str) -> str | None:
    stripped = line.strip()
    if not stripped:
        return None
    if stripped.startswith("#"):
        title = stripped.lstrip("#").strip()
        return title or None
    if re.match(r"^(第[一二三四五六七八九十\d]+[章节部分]|[一二三四五六七八九十\d]+[、.．])", stripped):
        return stripped
    if len(stripped) <= 32 and stripped.endswith(("：", ":")):
        return stripped.rstrip("：:")
    return None


@dataclass(frozen=True)
class TextChunk:
    index: int
    title_path: str
    content: str


def chunk_text(text: str, max_chars: int = 700, overlap_chars: int = 80) -> list[TextChunk]:
    cleaned = _clean_text(text)
    if not cleaned:
        return []

    chunks: list[TextChunk] = []
    current_title = ""
    current_parts: list[str] = []

    def flush() -> None:
        nonlocal current_parts
        content = "\n".join(part for part in current_parts if part.strip()).strip()
        current_parts = []
        if not content:
            return
        start = 0
        while start < len(content):
            end = min(len(content), start + max_chars)
            piece = content[start:end].strip()
            if piece:
                chunks.append(TextChunk(len(chunks), current_title, piece))
            if end >= len(content):
                break
            start = max(0, end - overlap_chars)

    for block in re.split(r"\n\s*\n", cleaned):
        block = block.strip()
        if not block:
            continue
        lines = [line.strip() for line in block.split("\n") if line.strip()]
        if len(lines) == 1:
            heading = _heading_title(lines[0])
            if heading:
                flush()
                current_title = heading
                continue
        if re.match(r"^第\d+行：", block):
            flush()
            current_parts.append(block)
            flush()
            continue
        projected = sum(len(part) for part in current_parts) + len(block)
        if current_parts and projected > max_chars:
            flush()
        current_parts.append(block)
    flush()
    return chunks


def _query_terms(query: str) -> list[str]:
    query = query.lower()
    terms: set[str] = set()
    for token in re.findall(r"[a-z0-9][a-z0-9_-]{1,}", query):
        terms.add(token)
    cjk = "".join(re.findall(r"[\u4e00-\u9fff]", query))
    for size in (2, 3, 4):
        for index in range(0, max(0, len(cjk) - size + 1)):
            terms.add(cjk[index : index + size])
    stop = {"这个", "那个", "可以", "怎么", "什么", "一下", "我们", "你们", "收到", "客户"}
    return sorted(term for term in terms if term and term not in stop)


def _keyword_score(query: str, title: str, content: str) -> float:
    terms = _query_terms(query)
    if not terms:
        return 0.0
    haystack = f"{title}\n{content}".lower()
    score = 0.0
    for term in terms:
        count = haystack.count(term.lower())
        if count <= 0:
            continue
        score += min(count, 4) * max(1.0, len(term) / 2.0)
        if term in title.lower():
            score += 2.0
    return score / max(1.0, len(terms))


def _fts_match_query(query: str, limit: int = 16) -> str:
    terms = _query_terms(query)
    quoted: list[str] = []
    for term in terms:
        term = term.strip()
        if len(term) < 2:
            continue
        escaped = term.replace('"', '""')
        quoted.append(f'"{escaped}"')
        if len(quoted) >= limit:
            break
    return " OR ".join(quoted)


def _fts_bm25_scores(conn: sqlite3.Connection, query: str, limit: int = 200) -> dict[str, float]:
    match_query = _fts_match_query(query)
    if not match_query:
        return {}
    try:
        rows = conn.execute(
            """
            SELECT chunk_id, bm25(knowledge_chunks_fts) AS rank
            FROM knowledge_chunks_fts
            WHERE knowledge_chunks_fts MATCH ?
            ORDER BY rank
            LIMIT ?
            """,
            (match_query, limit),
        ).fetchall()
    except sqlite3.Error:
        logging.debug("Knowledge FTS query failed query=%s match_query=%s", query, match_query, exc_info=True)
        return {}
    raw_scores: dict[str, float] = {}
    for row in rows:
        raw_scores[str(row["chunk_id"])] = max(0.0, -float(row["rank"] or 0.0))
    max_score = max(raw_scores.values(), default=0.0)
    if max_score <= 0:
        return {}
    return {chunk_id: score / max_score for chunk_id, score in raw_scores.items()}


def _model_terms(query: str) -> list[str]:
    terms = []
    for token in re.findall(r"\b[a-zA-Z]{1,8}[-_ ]?\d{1,4}[a-zA-Z0-9_-]*\b", query):
        terms.append(re.sub(r"\s+", "", token.lower()))
    return sorted(set(terms))


def _hybrid_weights(query: str) -> tuple[float, float, str]:
    text = query.lower()
    model_terms = _model_terms(query)
    precise_terms = [
        "实拍",
        "图片",
        "照片",
        "颜色",
        "尺寸",
        "规格",
        "配列",
        "型号",
        "轴体",
        "键帽",
        "k87",
        "k98",
        "k68",
    ]
    semantic_terms = [
        "发货",
        "当天",
        "售后",
        "保修",
        "退货",
        "换货",
        "质量",
        "安装",
        "怎么",
        "可以",
        "能不能",
    ]
    if model_terms or any(term in text for term in precise_terms):
        return 0.45, 0.55, "precise"
    if any(term in text for term in semantic_terms):
        return 0.65, 0.35, "semantic"
    return DEFAULT_VECTOR_WEIGHT, DEFAULT_KEYWORD_WEIGHT, "default"


def _metadata_boost(query: str, title: str = "", content: str = "", tags: list[str] | None = None, scenarios: list[str] | None = None, ocr_text: str = "") -> float:
    terms = _query_terms(query)
    model_terms = _model_terms(query)
    title_text = title.lower()
    content_text = content.lower()
    tags_text = " ".join(tags or []).lower()
    scenarios_text = " ".join(scenarios or []).lower()
    ocr_text = ocr_text.lower()
    boost = 0.0

    for term in model_terms:
        if term in title_text:
            boost += 0.22
        elif term in tags_text or term in scenarios_text:
            boost += 0.18
        elif term in ocr_text:
            boost += 0.14
        elif term in content_text:
            boost += 0.10

    title_hits = sum(1 for term in terms if term in title_text)
    tag_hits = sum(1 for term in terms if term in tags_text)
    scenario_hits = sum(1 for term in terms if term in scenarios_text)
    ocr_hits = sum(1 for term in terms if term in ocr_text)

    boost += min(0.18, title_hits * 0.04)
    boost += min(0.16, tag_hits * 0.04)
    boost += min(0.14, scenario_hits * 0.035)
    boost += min(0.12, ocr_hits * 0.03)
    return round(min(boost, 0.35), 4)


def _merge_adjacent_chunks(
    center: dict[str, Any],
    chunks_by_doc: dict[str, dict[int, dict[str, Any]]],
    used_chunk_ids: set[str],
    radius: int = ADJACENT_CHUNK_MERGE_RADIUS,
    max_chars: int = ADJACENT_CHUNK_MERGE_MAX_CHARS,
) -> dict[str, Any]:
    document_id = str(center.get("document_id") or "")
    center_index = int(center.get("chunk_index") or 0)
    if re.match(r"^第\d+行：", str(center.get("content") or "").strip()):
        content = str(center.get("content") or "").strip()
        return {
            "snippet": content,
            "chunk_ids": [str(center.get("chunk_id") or "")],
            "chunk_indexes": [center_index],
            "merged_count": 1,
        }
    doc_chunks = chunks_by_doc.get(document_id) or {}
    selected: dict[int, dict[str, Any]] = {center_index: center}

    for offset in range(1, max(0, radius) + 1):
        for index in (center_index - offset, center_index + offset):
            item = doc_chunks.get(index)
            if not item:
                continue
            chunk_id = str(item.get("chunk_id") or "")
            if chunk_id and chunk_id in used_chunk_ids:
                continue
            projected = "\n\n".join(str(value.get("content") or "") for _, value in sorted({**selected, index: item}.items()))
            if len(projected) > max_chars and index != center_index:
                continue
            selected[index] = item

    ordered = [value for _, value in sorted(selected.items())]
    snippet = "\n\n".join(str(value.get("content") or "").strip() for value in ordered if str(value.get("content") or "").strip())
    chunk_ids = [str(value.get("chunk_id") or "") for value in ordered if str(value.get("chunk_id") or "")]
    chunk_indexes = [int(value.get("chunk_index") or 0) for value in ordered]
    return {
        "snippet": snippet or str(center.get("content") or ""),
        "chunk_ids": chunk_ids,
        "chunk_indexes": chunk_indexes,
        "merged_count": len(chunk_ids),
    }


def _embedding_terms(text: str) -> list[str]:
    text = text.lower()
    terms: list[str] = []
    terms.extend(re.findall(r"[a-z0-9][a-z0-9_-]{1,}", text))
    cjk = "".join(re.findall(r"[\u4e00-\u9fff]", text))
    for size in (2, 3, 4):
        for index in range(0, max(0, len(cjk) - size + 1)):
            terms.append(cjk[index : index + size])
    return terms


def _stable_hash(value: str) -> int:
    digest = hashlib.blake2b(value.encode("utf-8"), digest_size=8).digest()
    return int.from_bytes(digest, "little", signed=False)


def _local_embedding(text: str, dim: int = LOCAL_EMBEDDING_DIM) -> list[float]:
    vector = [0.0] * dim
    for term in _embedding_terms(text):
        hashed = _stable_hash(term)
        index = hashed % dim
        sign = 1.0 if ((hashed >> 8) & 1) == 0 else -1.0
        vector[index] += sign * (1.0 + min(len(term), 8) * 0.08)
    norm = math.sqrt(sum(value * value for value in vector))
    if norm <= 0:
        return vector
    return [value / norm for value in vector]


def _pack_embedding(vector: list[float]) -> bytes:
    if not vector:
        return b""
    return struct.pack(f"<{len(vector)}f", *vector)


def _unpack_embedding(blob: bytes | memoryview | None) -> list[float]:
    if not blob:
        return []
    data = bytes(blob)
    if len(data) % 4 != 0:
        return []
    return list(struct.unpack(f"<{len(data) // 4}f", data))


def _cosine_similarity(left: list[float], right: list[float]) -> float:
    if not left or not right or len(left) != len(right):
        return 0.0
    return sum(a * b for a, b in zip(left, right))


def _configured_embedding_model() -> str:
    value = os.environ.get("AI_CUSTOMER_SERVICE_EMBEDDING_MODEL", "").strip()
    if value:
        return value
    if (BGE_SMALL_ZH_LOCAL_DIR / "modules.json").is_file():
        return str(BGE_SMALL_ZH_LOCAL_DIR)
    return BGE_SMALL_ZH_MODEL


def _normalize_embedding_model_name(model_name: str) -> str:
    value = (model_name or "").strip()
    if value:
        try:
            if Path(value).resolve() == BGE_SMALL_ZH_LOCAL_DIR.resolve():
                return BGE_SMALL_ZH_MODEL
        except OSError:
            pass
    if value in {"", "bge-small-zh-v1.5", "bge-small-zh"}:
        return BGE_SMALL_ZH_MODEL
    if value in {"local", "local-hash", LOCAL_EMBEDDING_MODEL}:
        return LOCAL_EMBEDDING_MODEL
    return value


def _configured_image_embedding_model() -> str:
    value = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_EMBEDDING_MODEL", "").strip()
    if value:
        return value
    if (BGE_VL_LOCAL_DIR / "config.json").is_file() or (BGE_VL_LOCAL_DIR / "preprocessor_config.json").is_file():
        return str(BGE_VL_LOCAL_DIR)
    if (CHINESE_CLIP_LOCAL_DIR / "config.json").is_file() or (CHINESE_CLIP_LOCAL_DIR / "preprocessor_config.json").is_file():
        return str(CHINESE_CLIP_LOCAL_DIR)
    if os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_EMBEDDING_ALLOW_DOWNLOAD", "").strip() == "1":
        return BGE_VL_MODEL
    return ""


def _normalize_image_embedding_model_name(model_name: str) -> str:
    value = (model_name or "").strip()
    if not value:
        return ""
    try:
        resolved = Path(value).resolve()
        if resolved == BGE_VL_LOCAL_DIR.resolve():
            return BGE_VL_MODEL
        if resolved == CHINESE_CLIP_LOCAL_DIR.resolve():
            return CHINESE_CLIP_MODEL
    except OSError:
        pass
    aliases = {
        "bge-vl": BGE_VL_MODEL,
        "bge-vl-base": BGE_VL_MODEL,
        "BAAI/bge-vl-base": BGE_VL_MODEL,
        "chinese-clip": CHINESE_CLIP_MODEL,
        "chinese-clip-vit-base": CHINESE_CLIP_MODEL,
        "OFA-Sys/chinese-clip-vit-base-patch16": CHINESE_CLIP_MODEL,
    }
    return aliases.get(value, value)


def _embedding_provider_for_model(model_name: str) -> EmbeddingProvider:
    normalized = _normalize_embedding_model_name(model_name)
    provider = _embedding_providers.get(normalized)
    if provider:
        return provider
    if normalized == LOCAL_EMBEDDING_MODEL:
        provider = LocalHashEmbeddingProvider()
    else:
        load_name = normalized
        if normalized == BGE_SMALL_ZH_MODEL and (BGE_SMALL_ZH_LOCAL_DIR / "modules.json").is_file():
            load_name = str(BGE_SMALL_ZH_LOCAL_DIR)
        provider = SentenceTransformerEmbeddingProvider(load_name, canonical_model_name=normalized)
    _embedding_providers[normalized] = provider
    return provider


def _image_embedding_load_name(normalized: str) -> str:
    if normalized == BGE_VL_MODEL and (
        (BGE_VL_LOCAL_DIR / "config.json").is_file()
        or (BGE_VL_LOCAL_DIR / "preprocessor_config.json").is_file()
    ):
        return str(BGE_VL_LOCAL_DIR)
    if normalized == CHINESE_CLIP_MODEL and (
        (CHINESE_CLIP_LOCAL_DIR / "config.json").is_file()
        or (CHINESE_CLIP_LOCAL_DIR / "preprocessor_config.json").is_file()
    ):
        return str(CHINESE_CLIP_LOCAL_DIR)
    return normalized


def _image_embedding_provider_for_model(model_name: str) -> ImageEmbeddingProvider:
    normalized = _normalize_image_embedding_model_name(model_name)
    if not normalized:
        raise EmbeddingProviderError("image embedding model is not configured")
    provider = _image_embedding_providers.get(normalized)
    if provider:
        return provider
    load_name = _image_embedding_load_name(normalized)
    if normalized == BGE_VL_MODEL:
        provider = SentenceTransformerImageEmbeddingProvider(
            load_name,
            canonical_model_name=normalized,
        )
    else:
        provider = TransformersImageEmbeddingProvider(
            load_name,
            canonical_model_name=normalized,
        )
    _image_embedding_providers[normalized] = provider
    return provider


def _active_embedding_provider() -> EmbeddingProvider:
    model_name = _normalize_embedding_model_name(_configured_embedding_model())
    try:
        return _embedding_provider_for_model(model_name)
    except EmbeddingProviderError:
        return _embedding_provider_for_model(LOCAL_EMBEDDING_MODEL)


def _encode_active_embedding(text: str, is_query: bool = False) -> tuple[str, int, list[float]]:
    provider = _active_embedding_provider()
    try:
        vector = provider.encode(text, is_query=is_query)
    except Exception:
        provider = _embedding_provider_for_model(LOCAL_EMBEDDING_MODEL)
        vector = provider.encode(text, is_query=is_query)
    return provider.model_name, len(vector) or provider.dim, vector


def _encode_active_image_embedding(path: Path) -> tuple[str, int, list[float]]:
    model_name = _normalize_image_embedding_model_name(_configured_image_embedding_model())
    if not model_name:
        return "", 0, []
    try:
        provider = _image_embedding_provider_for_model(model_name)
        vector = provider.encode_image(path)
        return provider.model_name, len(vector) or provider.dim, vector
    except Exception as exc:
        logging.warning("Knowledge image embedding disabled for path=%s model=%s error=%s", path, model_name, exc)
        return "", 0, []


def _encode_text_for_image_embedding_model(text: str, model_name: str) -> list[float]:
    normalized = _normalize_image_embedding_model_name(model_name)
    if not normalized:
        return []
    try:
        return _image_embedding_provider_for_model(normalized).encode_text(text)
    except Exception as exc:
        logging.warning("Knowledge image text embedding failed model=%s error=%s", normalized, exc)
        return []


def _encode_embedding_for_existing_model(text: str, model_name: str, is_query: bool = False) -> list[float]:
    if not (model_name or "").strip():
        return []
    normalized = _normalize_embedding_model_name(model_name)
    try:
        return _embedding_provider_for_model(normalized).encode(text, is_query=is_query)
    except EmbeddingProviderError:
        if normalized == LOCAL_EMBEDDING_MODEL:
            return _local_embedding(text)
        return []


def _configured_image_analyzer() -> ImageAssetAnalyzer:
    provider = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_ANALYZER", "").strip().lower()
    if provider not in {"ark", "doubao", "volcengine"}:
        return LocalImageAssetAnalyzer()
    api_key = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_ANALYSIS_API_KEY", "").strip()
    model = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_ANALYSIS_MODEL", "").strip()
    base_url = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_ANALYSIS_BASE_URL", "").strip()
    timeout_raw = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_ANALYSIS_TIMEOUT_SEC", "").strip()
    try:
        timeout = float(timeout_raw) if timeout_raw else 45.0
    except ValueError:
        timeout = 45.0
    try:
        return ArkDoubaoImageAssetAnalyzer(
            api_key=api_key,
            model=model,
            base_url=base_url or ARK_IMAGE_ANALYSIS_DEFAULT_BASE_URL,
            timeout=max(5.0, min(timeout, 180.0)),
        )
    except Exception:
        return LocalImageAssetAnalyzer()


class KnowledgeStore:
    def __init__(
        self,
        db_path: Path | None = None,
        files_dir: Path | None = None,
        connection_factory: Callable[[], sqlite3.Connection] | None = None,
        image_analyzer: ImageAssetAnalyzer | None = None,
    ) -> None:
        self.db_path = db_path or resolved_knowledge_db_path()
        self.files_dir = files_dir or resolved_knowledge_files_dir()
        self._connection_factory = connection_factory
        self._image_analyzer = image_analyzer or _configured_image_analyzer()
        self.ensure_schema()

    def _connect(self) -> sqlite3.Connection:
        if self._connection_factory:
            conn = self._connection_factory()
            conn.row_factory = sqlite3.Row
            return conn
        return _open_knowledge_db(self.db_path)

    def ensure_schema(self) -> None:
        conn = self._connect()
        try:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS knowledge_bases (
                    id TEXT PRIMARY KEY,
                    name TEXT NOT NULL,
                    description TEXT NOT NULL DEFAULT '',
                    enabled INTEGER NOT NULL DEFAULT 1,
                    platform TEXT NOT NULL DEFAULT '',
                    shop_id TEXT NOT NULL DEFAULT '',
                    robot_id TEXT NOT NULL DEFAULT '',
                    scene TEXT NOT NULL DEFAULT 'reply_draft',
                    source_directory TEXT NOT NULL DEFAULT '',
                    applicable_shops TEXT NOT NULL DEFAULT '',
                    last_import_status TEXT NOT NULL DEFAULT '',
                    last_import_error TEXT NOT NULL DEFAULT '',
                    last_import_elapsed_ms INTEGER NOT NULL DEFAULT 0,
                    last_import_document_count INTEGER NOT NULL DEFAULT 0,
                    last_import_image_count INTEGER NOT NULL DEFAULT 0,
                    last_import_at TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS knowledge_documents (
                    id TEXT PRIMARY KEY,
                    base_id TEXT NOT NULL,
                    title TEXT NOT NULL,
                    original_filename TEXT NOT NULL,
                    file_path TEXT NOT NULL,
                    file_type TEXT NOT NULL,
                    file_size INTEGER NOT NULL DEFAULT 0,
                    content_hash TEXT NOT NULL,
                    status TEXT NOT NULL,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    version INTEGER NOT NULL DEFAULT 1,
                    error_message TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS knowledge_chunks (
                    id TEXT PRIMARY KEY,
                    base_id TEXT NOT NULL,
                    document_id TEXT NOT NULL,
                    chunk_index INTEGER NOT NULL,
                    title_path TEXT NOT NULL DEFAULT '',
                    content TEXT NOT NULL,
                    token_estimate INTEGER NOT NULL DEFAULT 0,
                    content_hash TEXT NOT NULL,
                    embedding BLOB,
                    embedding_model TEXT NOT NULL DEFAULT '',
                    embedding_dim INTEGER NOT NULL DEFAULT 0,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE VIRTUAL TABLE IF NOT EXISTS knowledge_chunks_fts USING fts5(
                    chunk_id UNINDEXED,
                    base_id UNINDEXED,
                    document_id UNINDEXED,
                    title_path,
                    content
                );

                CREATE TABLE IF NOT EXISTS knowledge_retrieval_events (
                    id TEXT PRIMARY KEY,
                    query TEXT NOT NULL,
                    platform TEXT NOT NULL DEFAULT '',
                    shop_id TEXT NOT NULL DEFAULT '',
                    scene TEXT NOT NULL DEFAULT '',
                    result_count INTEGER NOT NULL DEFAULT 0,
                    latency_ms INTEGER NOT NULL DEFAULT 0,
                    mode TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS knowledge_image_assets (
                    id TEXT PRIMARY KEY,
                    base_id TEXT NOT NULL,
                    title TEXT NOT NULL,
                    original_filename TEXT NOT NULL,
                    file_path TEXT NOT NULL,
                    source_path TEXT NOT NULL DEFAULT '',
                    thumbnail_path TEXT NOT NULL DEFAULT '',
                    file_type TEXT NOT NULL,
                    file_size INTEGER NOT NULL DEFAULT 0,
                    content_hash TEXT NOT NULL,
                    width INTEGER NOT NULL DEFAULT 0,
                    height INTEGER NOT NULL DEFAULT 0,
                    asset_type TEXT NOT NULL DEFAULT 'other',
                    summary TEXT NOT NULL DEFAULT '',
                    ocr_text TEXT NOT NULL DEFAULT '',
                    tags_json TEXT NOT NULL DEFAULT '[]',
                    scenarios_json TEXT NOT NULL DEFAULT '[]',
                    risk_tags_json TEXT NOT NULL DEFAULT '[]',
                    suggested_reply TEXT NOT NULL DEFAULT '',
                    analysis_model TEXT NOT NULL DEFAULT '',
                    analysis_status TEXT NOT NULL DEFAULT 'pending',
                    raw_analysis_json TEXT NOT NULL DEFAULT '{}',
                    embedding BLOB,
                    embedding_model TEXT NOT NULL DEFAULT '',
                    embedding_dim INTEGER NOT NULL DEFAULT 0,
                    image_embedding BLOB,
                    image_embedding_model TEXT NOT NULL DEFAULT '',
                    image_embedding_dim INTEGER NOT NULL DEFAULT 0,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS knowledge_platform_bindings (
                    id TEXT PRIMARY KEY,
                    platform TEXT NOT NULL,
                    base_id TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    UNIQUE(platform, base_id)
                );
                """
            )
            self._ensure_column(conn, "knowledge_chunks", "embedding", "BLOB")
            self._ensure_column(conn, "knowledge_bases", "source_directory", "TEXT NOT NULL DEFAULT ''")
            self._ensure_column(conn, "knowledge_bases", "applicable_shops", "TEXT NOT NULL DEFAULT ''")
            self._ensure_column(conn, "knowledge_bases", "last_import_status", "TEXT NOT NULL DEFAULT ''")
            self._ensure_column(conn, "knowledge_bases", "last_import_error", "TEXT NOT NULL DEFAULT ''")
            self._ensure_column(conn, "knowledge_bases", "last_import_elapsed_ms", "INTEGER NOT NULL DEFAULT 0")
            self._ensure_column(conn, "knowledge_bases", "last_import_document_count", "INTEGER NOT NULL DEFAULT 0")
            self._ensure_column(conn, "knowledge_bases", "last_import_image_count", "INTEGER NOT NULL DEFAULT 0")
            self._ensure_column(conn, "knowledge_bases", "last_import_at", "TEXT NOT NULL DEFAULT ''")
            self._ensure_column(conn, "knowledge_chunks", "embedding_model", "TEXT NOT NULL DEFAULT ''")
            self._ensure_column(conn, "knowledge_chunks", "embedding_dim", "INTEGER NOT NULL DEFAULT 0")
            self._ensure_column(conn, "knowledge_image_assets", "thumbnail_path", "TEXT NOT NULL DEFAULT ''")
            self._ensure_column(conn, "knowledge_image_assets", "source_path", "TEXT NOT NULL DEFAULT ''")
            self._ensure_column(conn, "knowledge_image_assets", "width", "INTEGER NOT NULL DEFAULT 0")
            self._ensure_column(conn, "knowledge_image_assets", "height", "INTEGER NOT NULL DEFAULT 0")
            self._ensure_column(conn, "knowledge_image_assets", "embedding", "BLOB")
            self._ensure_column(conn, "knowledge_image_assets", "embedding_model", "TEXT NOT NULL DEFAULT ''")
            self._ensure_column(conn, "knowledge_image_assets", "embedding_dim", "INTEGER NOT NULL DEFAULT 0")
            self._ensure_column(conn, "knowledge_image_assets", "image_embedding", "BLOB")
            self._ensure_column(conn, "knowledge_image_assets", "image_embedding_model", "TEXT NOT NULL DEFAULT ''")
            self._ensure_column(conn, "knowledge_image_assets", "image_embedding_dim", "INTEGER NOT NULL DEFAULT 0")
            conn.commit()
        finally:
            conn.close()

    def _ensure_column(self, conn: sqlite3.Connection, table: str, column: str, definition: str) -> None:
        rows = conn.execute(f"PRAGMA table_info({table})").fetchall()
        if any(str(row["name"]) == column for row in rows):
            return
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")

    def create_base(
        self,
        name: str,
        description: str = "",
        platform: str = "",
        shop_id: str = "",
        robot_id: str = "",
        scene: str = "reply_draft",
        source_directory: str = "",
        applicable_shops: str = "",
        enabled: bool = True,
    ) -> dict[str, Any]:
        now = _now_text()
        base_id = _new_id("kb")
        conn = self._connect()
        try:
            conn.execute(
                """
                INSERT INTO knowledge_bases
                (id, name, description, enabled, platform, shop_id, robot_id, scene,
                 source_directory, applicable_shops, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    base_id,
                    name.strip() or "默认知识库",
                    description,
                    1 if enabled else 0,
                    platform,
                    shop_id,
                    robot_id,
                    scene,
                    source_directory,
                    applicable_shops,
                    now,
                    now,
                ),
            )
            conn.commit()
        finally:
            conn.close()
        return self.get_base(base_id) or {"id": base_id, "name": name}

    def update_base(
        self,
        base_id: str,
        *,
        name: str | None = None,
        description: str | None = None,
        source_directory: str | None = None,
        applicable_shops: str | None = None,
        enabled: bool | None = None,
        platform: str | None = None,
        shop_id: str | None = None,
        robot_id: str | None = None,
        scene: str | None = None,
    ) -> dict[str, Any]:
        if not base_id.strip():
            raise ValueError("missing knowledge base id")
        current = self.get_base(base_id)
        if not current:
            raise ValueError(f"knowledge base not found: {base_id}")

        updates: list[str] = []
        values: list[Any] = []
        fields: list[tuple[str, Any]] = [
            ("name", name.strip() if isinstance(name, str) else name),
            ("description", description),
            ("source_directory", source_directory),
            ("applicable_shops", applicable_shops),
            ("platform", platform),
            ("shop_id", shop_id),
            ("robot_id", robot_id),
            ("scene", scene),
        ]
        for column, value in fields:
            if value is None:
                continue
            if column == "name" and not str(value).strip():
                value = "默认知识库"
            updates.append(f"{column} = ?")
            values.append(str(value))
        if enabled is not None:
            updates.append("enabled = ?")
            values.append(1 if enabled else 0)
        updates.append("updated_at = ?")
        values.append(_now_text())
        values.append(base_id)

        conn = self._connect()
        try:
            conn.execute(f"UPDATE knowledge_bases SET {', '.join(updates)} WHERE id = ?", values)
            conn.commit()
        finally:
            conn.close()
        return self.get_base(base_id) or current

    def set_base_enabled(self, base_id: str, enabled: bool) -> dict[str, Any]:
        return self.update_base(base_id, enabled=enabled)

    def delete_base(self, base_id: str) -> dict[str, Any]:
        normalized_id = base_id.strip()
        if not normalized_id:
            raise ValueError("missing knowledge base id")
        current = self.get_base(normalized_id)
        if not current:
            raise ValueError(f"knowledge base not found: {normalized_id}")

        conn = self._connect()
        try:
            conn.execute("DELETE FROM knowledge_chunks_fts WHERE base_id = ?", (normalized_id,))
            conn.execute("DELETE FROM knowledge_chunks WHERE base_id = ?", (normalized_id,))
            conn.execute("DELETE FROM knowledge_documents WHERE base_id = ?", (normalized_id,))
            conn.execute("DELETE FROM knowledge_image_assets WHERE base_id = ?", (normalized_id,))
            conn.execute("DELETE FROM knowledge_platform_bindings WHERE base_id = ?", (normalized_id,))
            conn.execute("DELETE FROM knowledge_bases WHERE id = ?", (normalized_id,))
            conn.commit()
        finally:
            conn.close()
        logging.info("Knowledge base deleted base_id=%s name=%s", normalized_id, current.get("name", ""))
        return current

    def update_base_import_status(
        self,
        base_id: str,
        *,
        status: str,
        error: str = "",
        elapsed_ms: int = 0,
        document_count: int = 0,
        image_count: int = 0,
    ) -> dict[str, Any] | None:
        if not base_id.strip():
            return None
        now = _now_text()
        conn = self._connect()
        try:
            conn.execute(
                """
                UPDATE knowledge_bases
                SET last_import_status = ?,
                    last_import_error = ?,
                    last_import_elapsed_ms = ?,
                    last_import_document_count = ?,
                    last_import_image_count = ?,
                    last_import_at = ?,
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    status.strip(),
                    error.strip(),
                    max(0, int(elapsed_ms)),
                    max(0, int(document_count)),
                    max(0, int(image_count)),
                    now,
                    now,
                    base_id,
                ),
            )
            conn.commit()
        finally:
            conn.close()
        return self.get_base(base_id)

    def get_platform_bindings(self, platform: str) -> list[str]:
        normalized_platform = platform.strip().lower()
        if not normalized_platform:
            return []
        conn = self._connect()
        try:
            rows = conn.execute(
                """
                SELECT b.base_id
                FROM knowledge_platform_bindings b
                JOIN knowledge_bases kb ON kb.id = b.base_id
                WHERE b.platform = ?
                  AND kb.enabled = 1
                ORDER BY b.created_at ASC
                """,
                (normalized_platform,),
            ).fetchall()
            return [str(row["base_id"]) for row in rows]
        finally:
            conn.close()

    def set_platform_bindings(self, platform: str, base_ids: list[str]) -> list[str]:
        normalized_platform = platform.strip().lower()
        if not normalized_platform:
            raise ValueError("missing platform")
        normalized_base_ids: list[str] = []
        seen: set[str] = set()
        for base_id in base_ids:
            value = str(base_id or "").strip()
            if not value or value in seen:
                continue
            seen.add(value)
            normalized_base_ids.append(value)

        conn = self._connect()
        try:
            if normalized_base_ids:
                placeholders = ",".join("?" for _ in normalized_base_ids)
                rows = conn.execute(
                    f"SELECT id FROM knowledge_bases WHERE id IN ({placeholders})",
                    normalized_base_ids,
                ).fetchall()
                existing_ids = {str(row["id"]) for row in rows}
                missing = [base_id for base_id in normalized_base_ids if base_id not in existing_ids]
                if missing:
                    raise ValueError(f"knowledge base not found: {missing[0]}")

            now = _now_text()
            conn.execute("DELETE FROM knowledge_platform_bindings WHERE platform = ?", (normalized_platform,))
            for base_id in normalized_base_ids:
                conn.execute(
                    """
                    INSERT INTO knowledge_platform_bindings
                    (id, platform, base_id, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (_new_id("kbbind"), normalized_platform, base_id, now, now),
                )
            conn.commit()
        finally:
            conn.close()
        return self.get_platform_bindings(normalized_platform)

    def get_or_create_base(self, name: str, **kwargs: str) -> dict[str, Any]:
        conn = self._connect()
        try:
            row = conn.execute(
                "SELECT * FROM knowledge_bases WHERE name = ? AND enabled = 1 ORDER BY created_at LIMIT 1",
                (name.strip() or "默认知识库",),
            ).fetchone()
            if row:
                return _row_to_dict(conn, row)
        finally:
            conn.close()
        return self.create_base(name, **kwargs)

    def get_base(self, base_id: str) -> dict[str, Any] | None:
        conn = self._connect()
        try:
            row = conn.execute("SELECT * FROM knowledge_bases WHERE id = ?", (base_id,)).fetchone()
            return _row_to_dict(conn, row) if row else None
        finally:
            conn.close()

    def list_bases(self) -> list[dict[str, Any]]:
        conn = self._connect()
        try:
            rows = conn.execute(
                """
                SELECT b.*,
                       COALESCE(d.document_count, 0) AS document_count,
                       COALESCE(i.image_count, 0) AS image_count
                FROM knowledge_bases b
                LEFT JOIN (
                    SELECT base_id, COUNT(*) AS document_count
                    FROM knowledge_documents
                    WHERE enabled = 1
                    GROUP BY base_id
                ) d ON d.base_id = b.id
                LEFT JOIN (
                    SELECT base_id, COUNT(*) AS image_count
                    FROM knowledge_image_assets
                    WHERE enabled = 1
                    GROUP BY base_id
                ) i ON i.base_id = b.id
                ORDER BY b.created_at DESC
                """
            ).fetchall()
            return [_row_to_dict(conn, row) for row in rows]
        finally:
            conn.close()

    def import_file(self, base_id: str, source_path: Path, title: str | None = None) -> dict[str, Any]:
        source_path = source_path.resolve()
        if source_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"unsupported file type: {source_path.suffix}")
        if not source_path.is_file():
            raise FileNotFoundError(str(source_path))
        if not self.get_base(base_id):
            raise ValueError(f"knowledge base not found: {base_id}")

        data = source_path.read_bytes()
        content_hash = _sha256_bytes(data)
        now = _now_text()
        doc_title = title or source_path.stem

        existing_doc = self._find_existing_document(base_id, source_path.name, content_hash)
        if existing_doc:
            doc_id = str(existing_doc["id"])
            target_path = Path(str(existing_doc["file_path"]))
            if not target_path.name:
                target_path = self.files_dir / doc_id / source_path.name
            target_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source_path, target_path)
            conn = self._connect()
            try:
                conn.execute(
                    """
                    UPDATE knowledge_documents
                    SET title = ?, original_filename = ?, file_path = ?, file_type = ?,
                        file_size = ?, content_hash = ?, status = 'uploaded',
                        enabled = 1, error_message = '', updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        doc_title,
                        source_path.name,
                        str(target_path),
                        source_path.suffix.lower().lstrip("."),
                        len(data),
                        content_hash,
                        now,
                        doc_id,
                    ),
                )
                conn.commit()
            finally:
                conn.close()
            try:
                return self.reindex_document(doc_id)
            except Exception as exc:
                self._mark_document_failed(doc_id, str(exc))
                failed_doc = self.get_document(doc_id)
                if failed_doc:
                    return failed_doc
                raise

        doc_id = _new_id("doc")
        target_dir = self.files_dir / doc_id
        target_dir.mkdir(parents=True, exist_ok=True)
        target_path = target_dir / source_path.name
        shutil.copy2(source_path, target_path)

        conn = self._connect()
        try:
            conn.execute(
                """
                INSERT INTO knowledge_documents
                (id, base_id, title, original_filename, file_path, file_type, file_size,
                 content_hash, status, enabled, version, error_message, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'uploaded', 1, 1, '', ?, ?)
                """,
                (
                    doc_id,
                    base_id,
                    doc_title,
                    source_path.name,
                    str(target_path),
                    source_path.suffix.lower().lstrip("."),
                    len(data),
                    content_hash,
                    now,
                    now,
                ),
            )
            conn.commit()
        finally:
            conn.close()

        try:
            return self.reindex_document(doc_id)
        except Exception as exc:
            self._mark_document_failed(doc_id, str(exc))
            failed_doc = self.get_document(doc_id)
            if failed_doc:
                return failed_doc
            raise

    def _find_existing_document(self, base_id: str, original_filename: str, content_hash: str) -> dict[str, Any] | None:
        conn = self._connect()
        try:
            row = conn.execute(
                """
                SELECT * FROM knowledge_documents
                WHERE base_id = ?
                  AND original_filename = ?
                  AND content_hash = ?
                  AND enabled = 1
                ORDER BY created_at DESC
                LIMIT 1
                """,
                (base_id, original_filename, content_hash),
            ).fetchone()
            return _row_to_dict(conn, row) if row else None
        finally:
            conn.close()

    def import_image_file(self, base_id: str, source_path: Path, title: str | None = None) -> dict[str, Any]:
        source_path = source_path.resolve()
        started = time.monotonic()
        logging.info("Knowledge image import started path=%s base_id=%s", source_path, base_id)
        if source_path.suffix.lower() not in SUPPORTED_IMAGE_EXTENSIONS:
            raise ValueError(f"unsupported image type: {source_path.suffix}")
        if not source_path.is_file():
            raise FileNotFoundError(str(source_path))
        if not self.get_base(base_id):
            raise ValueError(f"knowledge base not found: {base_id}")

        data = source_path.read_bytes()
        content_hash = _sha256_bytes(data)
        now = _now_text()
        asset_title = title or _humanize_asset_name(source_path.stem)

        existing_asset = self._find_existing_image_asset(base_id, source_path.name, content_hash)
        asset_id = str(existing_asset["id"]) if existing_asset else _new_id("img")
        target_dir = self.files_dir / asset_id
        target_dir.mkdir(parents=True, exist_ok=True)
        target_path = target_dir / source_path.name
        shutil.copy2(source_path, target_path)

        try:
            logging.info("Knowledge image analysis started path=%s analyzer=%s", target_path, type(self._image_analyzer).__name__)
            analysis = self._image_analyzer.analyze(target_path)
            analysis_status = "ready"
            error_message = ""
            logging.info("Knowledge image analysis finished path=%s analyzer=%s", target_path, type(self._image_analyzer).__name__)
        except Exception as exc:
            logging.exception("Knowledge image analysis failed; falling back to local analyzer path=%s", target_path)
            analysis = LocalImageAssetAnalyzer().analyze(target_path)
            risk_tags = _string_list(analysis.get("risk_tags"), limit=16)
            if "analysis_failed" not in risk_tags:
                risk_tags.append("analysis_failed")
            analysis["risk_tags"] = risk_tags
            analysis["raw_analysis"] = {
                "source": "local_fallback_after_analysis_error",
                "fallback": analysis.get("raw_analysis") or {},
                "error": str(exc),
            }
            analysis_status = "needs_review"
            error_message = str(exc)

        width = int(analysis.get("width") or 0)
        height = int(analysis.get("height") or 0)
        if width <= 0 or height <= 0:
            width, height = _read_image_size(target_path)
        tags = _string_list(analysis.get("tags"), limit=16)
        scenarios = _string_list(analysis.get("scenarios"), limit=16)
        risk_tags = _string_list(analysis.get("risk_tags"), limit=16)
        summary = str(analysis.get("summary") or "").strip()
        ocr_text = str(analysis.get("ocr_text") or "").strip()
        suggested_reply = str(analysis.get("suggested_reply") or "").strip()
        asset_type = str(analysis.get("asset_type") or "other").strip() or "other"
        analysis_model = str(analysis.get("analysis_model") or "").strip()
        raw_analysis = analysis.get("raw_analysis")
        if not isinstance(raw_analysis, dict):
            raw_analysis = dict(analysis)
        if error_message:
            raw_analysis["error"] = error_message

        searchable_text = _image_asset_search_text(
            asset_title,
            asset_type,
            summary,
            ocr_text,
            tags,
            scenarios,
            risk_tags,
            suggested_reply,
        )
        embedding_model, embedding_dim, embedding_vector = _encode_active_embedding(searchable_text)
        embedding_blob = _pack_embedding(embedding_vector)
        image_embedding_model, image_embedding_dim, image_embedding_vector = _encode_active_image_embedding(target_path)
        image_embedding_blob = _pack_embedding(image_embedding_vector)

        conn = self._connect()
        try:
            if existing_asset:
                conn.execute(
                    """
                    UPDATE knowledge_image_assets
                    SET title = ?, original_filename = ?, file_path = ?, source_path = ?, file_type = ?,
                        file_size = ?, content_hash = ?, width = ?, height = ?,
                        asset_type = ?, summary = ?, ocr_text = ?, tags_json = ?,
                        scenarios_json = ?, risk_tags_json = ?, suggested_reply = ?,
                        analysis_model = ?, analysis_status = ?, raw_analysis_json = ?,
                        embedding = ?, embedding_model = ?, embedding_dim = ?,
                        image_embedding = ?, image_embedding_model = ?, image_embedding_dim = ?,
                        enabled = 1, updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        asset_title,
                        source_path.name,
                        str(target_path),
                        str(source_path),
                        source_path.suffix.lower().lstrip("."),
                        len(data),
                        content_hash,
                        width,
                        height,
                        asset_type,
                        summary,
                        ocr_text,
                        _json_dumps(tags),
                        _json_dumps(scenarios),
                        _json_dumps(risk_tags),
                        suggested_reply,
                        analysis_model,
                        analysis_status,
                        _json_dumps(raw_analysis),
                        embedding_blob,
                        embedding_model,
                        embedding_dim,
                        image_embedding_blob,
                        image_embedding_model,
                        image_embedding_dim,
                        now,
                        asset_id,
                    ),
                )
            else:
                conn.execute(
                    """
                    INSERT INTO knowledge_image_assets
                    (id, base_id, title, original_filename, file_path, source_path, thumbnail_path, file_type,
                     file_size, content_hash, width, height, asset_type, summary, ocr_text,
                     tags_json, scenarios_json, risk_tags_json, suggested_reply, analysis_model,
                     analysis_status, raw_analysis_json, embedding, embedding_model, embedding_dim,
                     image_embedding, image_embedding_model, image_embedding_dim,
                     enabled, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, '', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?)
                    """,
                    (
                        asset_id,
                        base_id,
                        asset_title,
                        source_path.name,
                        str(target_path),
                        str(source_path),
                        source_path.suffix.lower().lstrip("."),
                        len(data),
                        content_hash,
                        width,
                        height,
                        asset_type,
                        summary,
                        ocr_text,
                        _json_dumps(tags),
                        _json_dumps(scenarios),
                        _json_dumps(risk_tags),
                        suggested_reply,
                        analysis_model,
                        analysis_status,
                        _json_dumps(raw_analysis),
                        embedding_blob,
                        embedding_model,
                        embedding_dim,
                        image_embedding_blob,
                        image_embedding_model,
                        image_embedding_dim,
                        now,
                        now,
                    ),
                )
            conn.commit()
        finally:
            conn.close()
        result = self.get_image_asset(asset_id) or {"id": asset_id, "title": asset_title}
        logging.info(
            "Knowledge image import finished path=%s asset_id=%s status=%s elapsed_ms=%d",
            source_path,
            asset_id,
            analysis_status,
            int((time.monotonic() - started) * 1000),
        )
        return result

    def _find_existing_image_asset(self, base_id: str, original_filename: str, content_hash: str) -> dict[str, Any] | None:
        conn = self._connect()
        try:
            row = conn.execute(
                """
                SELECT * FROM knowledge_image_assets
                WHERE base_id = ?
                  AND original_filename = ?
                  AND content_hash = ?
                  AND enabled = 1
                ORDER BY created_at DESC
                LIMIT 1
                """,
                (base_id, original_filename, content_hash),
            ).fetchone()
            if row:
                return _image_asset_row_to_dict(row)
            row = conn.execute(
                """
                SELECT * FROM knowledge_image_assets
                WHERE base_id = ?
                  AND content_hash = ?
                  AND enabled = 1
                ORDER BY updated_at DESC, created_at DESC
                LIMIT 1
                """,
                (base_id, content_hash),
            ).fetchone()
            return _image_asset_row_to_dict(row) if row else None
        finally:
            conn.close()

    def import_directory(self, base_id: str, directory: Path) -> list[dict[str, Any]]:
        return self.import_directory_with_assets(base_id, directory)["documents"]

    def _sync_image_assets_for_directory(self, base_id: str, directory: Path, image_files: list[Path]) -> int:
        current_names = {path.name for path in image_files}
        current_paths = {str(path.resolve()) for path in image_files}
        now = _now_text()
        disabled = 0
        conn = self._connect()
        try:
            rows = conn.execute(
                """
                SELECT id, original_filename, source_path
                FROM knowledge_image_assets
                WHERE base_id = ? AND enabled = 1
                ORDER BY updated_at DESC, created_at DESC
                """,
                (base_id,),
            ).fetchall()
            for row in rows:
                source_path = str(row["source_path"] or "")
                original_filename = str(row["original_filename"] or "")
                keep = source_path in current_paths if source_path else original_filename in current_names
                if keep:
                    continue
                conn.execute(
                    "UPDATE knowledge_image_assets SET enabled = 0, updated_at = ? WHERE id = ?",
                    (now, row["id"]),
                )
                disabled += 1

            rows = conn.execute(
                """
                SELECT id, content_hash
                FROM knowledge_image_assets
                WHERE base_id = ? AND enabled = 1
                ORDER BY updated_at DESC, created_at DESC
                """,
                (base_id,),
            ).fetchall()
            seen_hashes: set[str] = set()
            for row in rows:
                content_hash = str(row["content_hash"] or "")
                if not content_hash or content_hash not in seen_hashes:
                    if content_hash:
                        seen_hashes.add(content_hash)
                    continue
                conn.execute(
                    "UPDATE knowledge_image_assets SET enabled = 0, updated_at = ? WHERE id = ?",
                    (now, row["id"]),
                )
                disabled += 1
            conn.commit()
            if disabled:
                logging.info(
                    "Knowledge image directory sync disabled stale assets directory=%s base_id=%s disabled=%d",
                    directory,
                    base_id,
                    disabled,
                )
            return disabled
        finally:
            conn.close()

    def import_directory_with_assets(self, base_id: str, directory: Path) -> dict[str, list[dict[str, Any]]]:
        started = time.monotonic()
        directory = directory.resolve()
        if not directory.is_dir():
            raise FileNotFoundError(str(directory))
        self.update_base(base_id, source_directory=str(directory))
        self.update_base_import_status(base_id, status="running")
        files = [
            file_path
            for file_path in sorted(directory.rglob("*"))
            if file_path.is_file() and not file_path.name.startswith("~$")
        ]
        supported_docs = [file_path for file_path in files if file_path.suffix.lower() in SUPPORTED_EXTENSIONS]
        supported_images = [file_path for file_path in files if file_path.suffix.lower() in SUPPORTED_IMAGE_EXTENSIONS]
        logging.info(
            "Knowledge directory import started directory=%s base_id=%s files=%d docs=%d images=%d",
            directory,
            base_id,
            len(files),
            len(supported_docs),
            len(supported_images),
        )
        imported: list[dict[str, Any]] = []
        imported_images: list[dict[str, Any]] = []
        for file_path in files:
            if file_path.suffix.lower() in SUPPORTED_EXTENSIONS and file_path.is_file():
                try:
                    logging.info("Knowledge document import started path=%s", file_path)
                    imported.append(self.import_file(base_id, file_path))
                    logging.info("Knowledge document import finished path=%s status=%s", file_path, imported[-1].get("status"))
                except Exception as exc:
                    logging.exception("Knowledge document import failed path=%s", file_path)
                    imported.append(
                        {
                            "id": "",
                            "base_id": base_id,
                            "title": file_path.stem,
                            "original_filename": file_path.name,
                            "file_path": str(file_path),
                            "file_type": file_path.suffix.lower().lstrip("."),
                            "file_size": file_path.stat().st_size if file_path.exists() else 0,
                            "status": "failed",
                            "enabled": 0,
                            "version": 1,
                            "error_message": str(exc),
                            "created_at": _now_text(),
                            "updated_at": _now_text(),
                        }
                    )
            elif file_path.suffix.lower() in SUPPORTED_IMAGE_EXTENSIONS and file_path.is_file():
                try:
                    imported_images.append(self.import_image_file(base_id, file_path))
                except Exception as exc:
                    logging.exception("Knowledge image import failed path=%s", file_path)
                    imported_images.append(
                        {
                            "id": "",
                            "base_id": base_id,
                            "title": file_path.stem,
                            "original_filename": file_path.name,
                            "file_path": str(file_path),
                            "file_type": file_path.suffix.lower().lstrip("."),
                            "file_size": file_path.stat().st_size if file_path.exists() else 0,
                            "asset_type": "other",
                            "summary": "",
                            "ocr_text": "",
                            "tags": [],
                            "scenarios": [],
                            "risk_tags": ["import_failed"],
                            "analysis_status": "error",
                            "error_message": str(exc),
                            "created_at": _now_text(),
                            "updated_at": _now_text(),
                        }
                    )
        self._sync_image_assets_for_directory(base_id, directory, supported_images)
        elapsed_ms = int((time.monotonic() - started) * 1000)
        failed_documents = [item for item in imported if item.get("status") == "failed"]
        failed_images = [item for item in imported_images if item.get("analysis_status") == "error"]
        import_status = "success" if not failed_documents and not failed_images else "partial"
        error_parts = []
        if failed_documents:
            error_parts.append(f"documents_failed={len(failed_documents)}")
        if failed_images:
            error_parts.append(f"images_failed={len(failed_images)}")
        self.update_base_import_status(
            base_id,
            status=import_status,
            error="; ".join(error_parts),
            elapsed_ms=elapsed_ms,
            document_count=len(imported),
            image_count=len(imported_images),
        )
        logging.info(
            "Knowledge directory import finished directory=%s base_id=%s documents=%d images=%d elapsed_ms=%d",
            directory,
            base_id,
            len(imported),
            len(imported_images),
            elapsed_ms,
        )
        return {"documents": imported, "images": imported_images}

    def reindex_document(self, document_id: str) -> dict[str, Any]:
        doc = self.get_document(document_id)
        if not doc:
            raise ValueError(f"document not found: {document_id}")

        now = _now_text()
        conn = self._connect()
        try:
            conn.execute(
                "UPDATE knowledge_documents SET status = 'extracting', error_message = '', updated_at = ? WHERE id = ?",
                (now, document_id),
            )
            conn.commit()
        finally:
            conn.close()

        text = extract_text(Path(doc["file_path"]))
        chunks = chunk_text(text)
        if not chunks:
            raise RuntimeError("no extractable text chunks")

        conn = self._connect()
        try:
            conn.execute(
                "UPDATE knowledge_documents SET status = 'indexing', updated_at = ? WHERE id = ?",
                (now, document_id),
            )
            conn.execute("DELETE FROM knowledge_chunks WHERE document_id = ?", (document_id,))
            conn.execute("DELETE FROM knowledge_chunks_fts WHERE document_id = ?", (document_id,))
            for chunk in chunks:
                chunk_id = _new_id("chunk")
                content_hash = hashlib.sha256(chunk.content.encode("utf-8")).hexdigest()
                token_estimate = max(1, len(chunk.content) // 2)
                embedding_text = f"{doc['title']}\n{chunk.title_path}\n{chunk.content}"
                embedding_model, embedding_dim, embedding_vector = _encode_active_embedding(embedding_text)
                embedding_blob = _pack_embedding(embedding_vector)
                conn.execute(
                    """
                    INSERT INTO knowledge_chunks
                    (id, base_id, document_id, chunk_index, title_path, content,
                     token_estimate, content_hash, embedding, embedding_model, embedding_dim,
                     enabled, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?)
                    """,
                    (
                        chunk_id,
                        doc["base_id"],
                        document_id,
                        chunk.index,
                        chunk.title_path,
                        chunk.content,
                        token_estimate,
                        content_hash,
                        embedding_blob,
                        embedding_model,
                        embedding_dim,
                        now,
                        now,
                    ),
                )
                conn.execute(
                    """
                    INSERT INTO knowledge_chunks_fts
                    (chunk_id, base_id, document_id, title_path, content)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (chunk_id, doc["base_id"], document_id, chunk.title_path, chunk.content),
                )
            conn.execute(
                "UPDATE knowledge_documents SET status = 'ready', updated_at = ? WHERE id = ?",
                (now, document_id),
            )
            conn.commit()
        finally:
            conn.close()
        return self.get_document(document_id) or doc

    def _mark_document_failed(self, document_id: str, message: str) -> None:
        conn = self._connect()
        try:
            conn.execute(
                "UPDATE knowledge_documents SET status = 'failed', error_message = ?, updated_at = ? WHERE id = ?",
                (message[:1000], _now_text(), document_id),
            )
            conn.commit()
        finally:
            conn.close()

    def get_document(self, document_id: str) -> dict[str, Any] | None:
        conn = self._connect()
        try:
            row = conn.execute("SELECT * FROM knowledge_documents WHERE id = ?", (document_id,)).fetchone()
            return _row_to_dict(conn, row) if row else None
        finally:
            conn.close()

    def list_documents(self, base_id: str | None = None) -> list[dict[str, Any]]:
        conn = self._connect()
        try:
            if base_id:
                rows = conn.execute(
                    """
                    SELECT d.*,
                           COALESCE(COUNT(c.id), 0) AS chunk_count
                    FROM knowledge_documents d
                    LEFT JOIN knowledge_chunks c ON c.document_id = d.id AND c.enabled = 1
                    WHERE d.base_id = ?
                      AND d.enabled = 1
                    GROUP BY d.id
                    ORDER BY d.created_at DESC
                    """,
                    (base_id,),
                ).fetchall()
            else:
                rows = conn.execute(
                    """
                    SELECT d.*,
                           COALESCE(COUNT(c.id), 0) AS chunk_count
                    FROM knowledge_documents d
                    LEFT JOIN knowledge_chunks c ON c.document_id = d.id AND c.enabled = 1
                    WHERE d.enabled = 1
                    GROUP BY d.id
                    ORDER BY d.created_at DESC
                    """
                ).fetchall()
            return [_row_to_dict(conn, row) for row in rows]
        finally:
            conn.close()

    def get_image_asset(self, asset_id: str) -> dict[str, Any] | None:
        conn = self._connect()
        try:
            row = conn.execute("SELECT * FROM knowledge_image_assets WHERE id = ?", (asset_id,)).fetchone()
            return _image_asset_row_to_dict(row) if row else None
        finally:
            conn.close()

    def list_image_assets(self, base_id: str | None = None) -> list[dict[str, Any]]:
        conn = self._connect()
        try:
            if base_id:
                rows = conn.execute(
                    """
                    SELECT *
                    FROM knowledge_image_assets
                    WHERE base_id = ?
                      AND enabled = 1
                    ORDER BY created_at DESC
                    """,
                    (base_id,),
                ).fetchall()
            else:
                rows = conn.execute(
                    """
                    SELECT *
                    FROM knowledge_image_assets
                    WHERE enabled = 1
                    ORDER BY created_at DESC
                    """
                ).fetchall()
            return [_image_asset_row_to_dict(row) for row in rows]
        finally:
            conn.close()

    def search_image_assets(
        self,
        query: str,
        base_ids: list[str] | None = None,
        platform: str = "",
        shop_id: str = "",
        scene: str = "",
        top_k: int = 5,
        mode: str = "hybrid",
        include_risky: bool = False,
    ) -> dict[str, Any]:
        started = time.perf_counter()
        top_k = max(1, min(int(top_k or 5), 20))
        query = query.strip()
        mode = (mode or "hybrid").strip().lower()
        if mode not in {"keyword", "vector", "hybrid"}:
            mode = "hybrid"
        if not query:
            return {"status": "success", "results": [], "metadata": {"mode": mode, "latency_ms": 0}}

        conn = self._connect()
        try:
            rows = conn.execute(
                """
                SELECT a.*,
                       b.platform,
                       b.shop_id,
                       b.scene
                FROM knowledge_image_assets a
                JOIN knowledge_bases b ON b.id = a.base_id
                WHERE a.enabled = 1
                  AND a.analysis_status IN ('ready', 'needs_review')
                  AND b.enabled = 1
                """
            ).fetchall()
        finally:
            conn.close()

        query_embeddings: dict[str, list[float]] = {}
        image_query_embeddings: dict[str, list[float]] = {}
        candidates: list[dict[str, Any]] = []
        base_filter = set(base_ids or [])
        normalized_query = query.lower()
        filename_query = bool(re.search(r"\.(?:jpg|jpeg|png|webp|bmp|gif)\b", query, re.IGNORECASE))
        for row in rows:
            item = _image_asset_row_to_dict(row)
            item["platform"] = row["platform"]
            item["shop_id"] = row["shop_id"]
            item["scene"] = row["scene"]
            if base_filter and item["base_id"] not in base_filter:
                continue
            if platform and item["platform"] and item["platform"] != platform:
                continue
            if shop_id and item["shop_id"] and item["shop_id"] != shop_id:
                continue
            if scene and item["scene"] and item["scene"] != scene:
                continue
            risk_tags = item.get("risk_tags") or []
            if risk_tags and not include_risky:
                continue
            if _explicit_model_mismatch(query, item):
                continue
            original_filename = str(item.get("original_filename") or "").lower()
            title_text = str(item.get("title") or "").lower()
            exact_filename_match = 0.0
            if original_filename and original_filename in normalized_query:
                exact_filename_match = 1.0
            elif title_text and title_text in normalized_query:
                exact_filename_match = 0.9
            if filename_query and exact_filename_match <= 0:
                continue

            searchable_text = _image_asset_search_text(
                item["title"],
                item["asset_type"],
                item["summary"],
                item["ocr_text"],
                item["tags"],
                item["scenarios"],
                risk_tags,
                item["suggested_reply"],
            )
            keyword_score = _keyword_score(query, f'{item["title"]} {item["asset_type"]}', searchable_text)
            metadata_boost = _metadata_boost(
                query,
                title=f'{item["title"]} {item["asset_type"]}',
                content=item["summary"],
                tags=item["tags"],
                scenarios=item["scenarios"],
                ocr_text=item["ocr_text"],
            )
            vector_score = 0.0
            embedding_model = str(item.get("embedding_model") or "")
            if exact_filename_match <= 0 and mode in {"vector", "hybrid"} and embedding_model:
                if embedding_model not in query_embeddings:
                    query_embeddings[embedding_model] = _encode_embedding_for_existing_model(query, embedding_model, is_query=True)
                vector_score = _cosine_similarity(
                    query_embeddings.get(embedding_model) or [],
                    _unpack_embedding(row["embedding"]),
                )
            image_vector_score = 0.0
            image_embedding_model = str(item.get("image_embedding_model") or "")
            if exact_filename_match <= 0 and mode in {"vector", "hybrid"} and image_embedding_model:
                if image_embedding_model not in image_query_embeddings:
                    image_query_embeddings[image_embedding_model] = _encode_text_for_image_embedding_model(query, image_embedding_model)
                image_vector_score = _cosine_similarity(
                    image_query_embeddings.get(image_embedding_model) or [],
                    _unpack_embedding(row["image_embedding"]),
                )
            best_vector_score = max(vector_score, image_vector_score)

            if mode == "keyword" and keyword_score <= 0:
                continue
            if mode == "vector" and best_vector_score < VECTOR_SCORE_THRESHOLD:
                continue
            if mode == "hybrid" and keyword_score <= 0 and best_vector_score < VECTOR_SCORE_THRESHOLD:
                continue
            item["keyword_score"] = keyword_score
            item["vector_score"] = vector_score
            item["image_vector_score"] = image_vector_score
            item["metadata_boost"] = metadata_boost
            item["rule_bonus"] = _image_rule_bonus(query, item)
            item["exact_filename_match"] = exact_filename_match
            candidates.append(item)

        max_keyword_score = max((float(item["keyword_score"]) for item in candidates), default=0.0)
        vector_weight, keyword_weight, weight_reason = _hybrid_weights(query)
        for item in candidates:
            keyword_norm = float(item["keyword_score"]) / max_keyword_score if max_keyword_score > 0 else 0.0
            vector_score = max(0.0, float(item["vector_score"]))
            image_vector_score = max(0.0, float(item.get("image_vector_score") or 0.0))
            rule_bonus = max(0.0, min(1.0, float(item.get("rule_bonus") or 0.0)))
            exact_filename_match = float(item.get("exact_filename_match") or 0.0)
            if mode == "keyword":
                score = keyword_norm + float(item.get("metadata_boost") or 0.0) + rule_bonus
                match_type = "keyword"
            elif mode == "vector":
                score = max(vector_score, image_vector_score) + float(item.get("metadata_boost") or 0.0) * 0.5 + rule_bonus
                match_type = "image_vector" if image_vector_score >= vector_score else "vector"
            else:
                text_score = min(1.0, max(keyword_norm, vector_score) + float(item.get("metadata_boost") or 0.0) * 0.3)
                score = (
                    DEFAULT_IMAGE_VECTOR_WEIGHT * image_vector_score
                    + DEFAULT_IMAGE_TEXT_WEIGHT * text_score
                    + DEFAULT_IMAGE_RULE_WEIGHT * rule_bonus
                )
                if image_vector_score >= VECTOR_SCORE_THRESHOLD and (keyword_norm > 0 or vector_score >= VECTOR_SCORE_THRESHOLD):
                    match_type = "hybrid_image"
                elif image_vector_score >= VECTOR_SCORE_THRESHOLD:
                    match_type = "image_vector"
                elif vector_score >= VECTOR_SCORE_THRESHOLD and keyword_norm > 0:
                    match_type = "hybrid"
                elif vector_score >= VECTOR_SCORE_THRESHOLD:
                    match_type = "text_vector"
                else:
                    match_type = "keyword"
            if exact_filename_match > 0:
                score += 2.0 * exact_filename_match
                match_type = "exact_filename"
            if risk_tags := item.get("risk_tags"):
                score -= min(0.4, 0.08 * len(risk_tags))
            if item.get("analysis_status") != "ready":
                score *= 0.6
            item["score"] = round(score, 4)
            item["match_type"] = match_type
            item["keyword_score"] = round(float(item["keyword_score"]), 4)
            item["vector_score"] = round(vector_score, 4)
            item["image_vector_score"] = round(image_vector_score, 4)
            item["metadata_boost"] = round(float(item.get("metadata_boost") or 0.0), 4)
            item["rule_bonus"] = round(rule_bonus, 4)
            item["exact_filename_match"] = exact_filename_match

        candidates.sort(key=lambda value: value["score"], reverse=True)
        should_rerank = not filename_query and not any(
            float(item.get("exact_filename_match") or 0.0) > 0 for item in candidates[:top_k]
        )
        rerank_metadata = _rerank_image_candidates(query, candidates[: min(5, max(top_k, 3))]) if should_rerank else {}
        if rerank_metadata:
            for index, adjustment in rerank_metadata.items():
                if 0 <= index < len(candidates):
                    candidates[index]["score"] = round(float(candidates[index]["score"]) + float(adjustment.get("score_boost") or 0.0), 4)
                    candidates[index]["rerank_score"] = round(float(adjustment.get("rerank_score") or 0.0), 4)
                    candidates[index]["rerank_reason"] = str(adjustment.get("reason") or "").strip()
                    if str(adjustment.get("match_type") or "").strip():
                        candidates[index]["match_type"] = str(adjustment["match_type"])
            candidates.sort(key=lambda value: value["score"], reverse=True)
        results: list[dict[str, Any]] = []
        for item in candidates[:top_k]:
            results.append(
                {
                    "asset_id": item["id"],
                    "source_id": item["id"],
                    "source_title": item["title"],
                    "original_filename": item["original_filename"],
                    "file_path": item["file_path"],
                    "thumbnail_path": item["thumbnail_path"],
                    "asset_type": item["asset_type"],
                    "summary": item["summary"],
                    "ocr_text": item["ocr_text"],
                    "tags": item["tags"],
                    "scenarios": item["scenarios"],
                    "risk_tags": item["risk_tags"],
                    "suggested_reply": item["suggested_reply"],
                    "score": item["score"],
                    "match_type": item["match_type"],
                    "recommendation": {
                        "should_attach": not bool(item["risk_tags"]) and item["analysis_status"] == "ready",
                        "reason": _image_recommendation_reason(query, item),
                        "requires_human_confirm": True,
                        "risk_notice": " / ".join(item["risk_tags"]) if item["risk_tags"] else "",
                    },
                    "metadata": {
                        "base_id": item["base_id"],
                        "platform": item["platform"],
                        "shop_id": item["shop_id"],
                        "scene": item["scene"],
                        "width": item["width"],
                        "height": item["height"],
                        "analysis_model": item["analysis_model"],
                        "analysis_status": item["analysis_status"],
                        "keyword_score": item["keyword_score"],
                        "vector_score": item["vector_score"],
                        "text_vector_score": item["vector_score"],
                        "image_vector_score": item["image_vector_score"],
                        "metadata_boost": item["metadata_boost"],
                        "rule_bonus": item["rule_bonus"],
                        "rerank_score": item.get("rerank_score", 0.0),
                        "rerank_reason": item.get("rerank_reason", ""),
                        "exact_filename_match": item["exact_filename_match"],
                        "vector_weight": vector_weight,
                        "keyword_weight": keyword_weight,
                        "image_vector_weight": DEFAULT_IMAGE_VECTOR_WEIGHT,
                        "image_text_weight": DEFAULT_IMAGE_TEXT_WEIGHT,
                        "image_rule_weight": DEFAULT_IMAGE_RULE_WEIGHT,
                        "weight_reason": weight_reason,
                        "embedding_model": item.get("embedding_model") or "",
                        "image_embedding_model": item.get("image_embedding_model") or "",
                    },
                }
            )

        latency_ms = int((time.perf_counter() - started) * 1000)
        return {"status": "success", "results": results, "metadata": {"mode": mode, "latency_ms": latency_ms}}

    def delete_document(self, document_id: str) -> bool:
        conn = self._connect()
        try:
            doc = conn.execute("SELECT * FROM knowledge_documents WHERE id = ?", (document_id,)).fetchone()
            if not doc:
                return False
            conn.execute("UPDATE knowledge_documents SET status = 'deleted', enabled = 0, updated_at = ? WHERE id = ?", (_now_text(), document_id))
            conn.execute("UPDATE knowledge_chunks SET enabled = 0, updated_at = ? WHERE document_id = ?", (_now_text(), document_id))
            conn.execute("DELETE FROM knowledge_chunks_fts WHERE document_id = ?", (document_id,))
            conn.commit()
            return True
        finally:
            conn.close()

    def search(
        self,
        query: str,
        base_ids: list[str] | None = None,
        platform: str = "",
        shop_id: str = "",
        scene: str = "",
        top_k: int = 5,
        mode: str = "hybrid",
    ) -> dict[str, Any]:
        started = time.perf_counter()
        top_k = max(1, min(int(top_k or 5), 20))
        query = query.strip()
        mode = (mode or "hybrid").strip().lower()
        if mode not in {"keyword", "vector", "hybrid"}:
            mode = "hybrid"
        if not query:
            return {"status": "success", "results": [], "metadata": {"mode": mode, "latency_ms": 0}}
        query_embeddings: dict[str, list[float]] = {}

        conn = self._connect()
        try:
            fts_scores = _fts_bm25_scores(conn, query)
            rows = conn.execute(
                """
                SELECT c.id AS chunk_id, c.base_id, c.document_id, c.chunk_index,
                       c.title_path, c.content, c.embedding, c.embedding_model, c.embedding_dim,
                       d.title AS source_title,
                       b.platform, b.shop_id, b.scene
                FROM knowledge_chunks c
                JOIN knowledge_documents d ON d.id = c.document_id
                JOIN knowledge_bases b ON b.id = c.base_id
                WHERE c.enabled = 1
                  AND d.enabled = 1
                  AND d.status = 'ready'
                  AND b.enabled = 1
                """
            ).fetchall()
        finally:
            conn.close()

        candidates: list[dict[str, Any]] = []
        chunks_by_doc: dict[str, dict[int, dict[str, Any]]] = {}
        base_filter = set(base_ids or [])
        for row in rows:
            item = dict(row)
            if base_filter and item["base_id"] not in base_filter:
                continue
            if platform and item["platform"] and item["platform"] != platform:
                continue
            if shop_id and item["shop_id"] and item["shop_id"] != shop_id:
                continue
            if scene and item["scene"] and item["scene"] != scene:
                continue
            chunks_by_doc.setdefault(str(item["document_id"]), {})[int(item["chunk_index"])] = item
            fts_score = fts_scores.get(str(item["chunk_id"]), 0.0)
            keyword_score = _keyword_score(query, f'{item["source_title"]} {item["title_path"]}', item["content"])
            if fts_score > 0:
                keyword_score += 2.0 * fts_score
            vector_score = 0.0
            embedding_model = str(item.get("embedding_model") or "")
            if mode in {"vector", "hybrid"} and embedding_model:
                if embedding_model not in query_embeddings:
                    query_embeddings[embedding_model] = _encode_embedding_for_existing_model(query, embedding_model, is_query=True)
                query_embedding = query_embeddings.get(embedding_model) or []
                vector_score = _cosine_similarity(query_embedding, _unpack_embedding(item.get("embedding")))

            if mode == "keyword" and keyword_score <= 0:
                continue
            if mode == "vector" and vector_score < VECTOR_SCORE_THRESHOLD:
                continue
            if mode == "hybrid" and keyword_score <= 0 and vector_score < VECTOR_SCORE_THRESHOLD:
                continue
            item["keyword_score"] = keyword_score
            item["vector_score"] = vector_score
            item["fts_score"] = fts_score
            item["metadata_boost"] = _metadata_boost(
                query,
                title=f'{item["source_title"]} {item["title_path"]}',
                content=item["content"],
            )
            candidates.append(item)

        max_keyword_score = max((float(item["keyword_score"]) for item in candidates), default=0.0)
        vector_weight, keyword_weight, weight_reason = _hybrid_weights(query)
        for item in candidates:
            keyword_norm = float(item["keyword_score"]) / max_keyword_score if max_keyword_score > 0 else 0.0
            vector_score = max(0.0, float(item["vector_score"]))
            if mode == "keyword":
                score = keyword_norm + float(item.get("metadata_boost") or 0.0)
                match_type = "keyword"
            elif mode == "vector":
                score = vector_score + float(item.get("metadata_boost") or 0.0) * 0.5
                match_type = "vector"
            else:
                score = vector_weight * vector_score + keyword_weight * keyword_norm + float(item.get("metadata_boost") or 0.0)
                if vector_score >= VECTOR_SCORE_THRESHOLD and keyword_norm > 0:
                    match_type = "hybrid"
                elif vector_score >= VECTOR_SCORE_THRESHOLD:
                    match_type = "vector"
                else:
                    match_type = "keyword"
            item["score"] = round(score, 4)
            item["match_type"] = match_type
            item["keyword_score"] = round(float(item["keyword_score"]), 4)
            item["vector_score"] = round(vector_score, 4)
            item["fts_score"] = round(float(item.get("fts_score") or 0.0), 4)
            item["metadata_boost"] = round(float(item.get("metadata_boost") or 0.0), 4)

        candidates.sort(key=lambda value: value["score"], reverse=True)
        results: list[dict[str, Any]] = []
        per_doc: dict[str, int] = {}
        used_chunk_ids: set[str] = set()
        for item in candidates:
            if str(item["chunk_id"]) in used_chunk_ids:
                continue
            if per_doc.get(item["document_id"], 0) >= 2:
                continue
            per_doc[item["document_id"]] = per_doc.get(item["document_id"], 0) + 1
            merged = _merge_adjacent_chunks(item, chunks_by_doc, used_chunk_ids)
            used_chunk_ids.update(merged["chunk_ids"])
            results.append(
                {
                    "source_id": item["document_id"],
                    "chunk_id": item["chunk_id"],
                    "source_title": item["source_title"],
                    "title_path": item["title_path"],
                    "snippet": merged["snippet"],
                    "score": item["score"],
                    "match_type": item["match_type"],
                    "metadata": {
                        "base_id": item["base_id"],
                        "platform": item["platform"],
                        "shop_id": item["shop_id"],
                        "scene": item["scene"],
                        "keyword_score": item["keyword_score"],
                        "vector_score": item["vector_score"],
                        "fts_score": item["fts_score"],
                        "metadata_boost": item["metadata_boost"],
                        "vector_weight": vector_weight,
                        "keyword_weight": keyword_weight,
                        "weight_reason": weight_reason,
                        "merged_chunk_ids": merged["chunk_ids"],
                        "merged_chunk_indexes": merged["chunk_indexes"],
                        "merged_chunk_count": merged["merged_count"],
                        "embedding_model": item.get("embedding_model") or "",
                    },
                }
            )
            if len(results) >= top_k:
                break

        latency_ms = int((time.perf_counter() - started) * 1000)
        self._record_retrieval(query, platform, shop_id, scene, len(results), latency_ms, mode)
        return {"status": "success", "results": results, "metadata": {"mode": mode, "latency_ms": latency_ms}}

    def warm_up_embeddings(self, text: str = "键盘 客服 检索 预热") -> dict[str, Any]:
        started = time.perf_counter()
        models: list[str] = []
        errors: list[str] = []

        conn = self._connect()
        try:
            rows = conn.execute(
                """
                SELECT DISTINCT embedding_model
                FROM knowledge_chunks
                WHERE enabled = 1
                  AND embedding_model IS NOT NULL
                  AND TRIM(embedding_model) != ''
                """
            ).fetchall()
        finally:
            conn.close()

        model_names = [str(row["embedding_model"] or "").strip() for row in rows]
        if not model_names:
            model_names = [_normalize_embedding_model_name(_configured_embedding_model())]

        for model_name in sorted(set(model_names)):
            if not model_name:
                continue
            try:
                vector = _encode_embedding_for_existing_model(text, model_name, is_query=True)
                if not vector:
                    raise EmbeddingProviderError(f"empty embedding for {model_name}")
                models.append(_normalize_embedding_model_name(model_name))
            except Exception as exc:  # pragma: no cover - defensive startup warmup
                errors.append(f"{model_name}: {exc}")

        image_model_names: list[str] = []
        conn = self._connect()
        try:
            rows = conn.execute(
                """
                SELECT DISTINCT image_embedding_model
                FROM knowledge_image_assets
                WHERE enabled = 1
                  AND image_embedding_model IS NOT NULL
                  AND TRIM(image_embedding_model) != ''
                """
            ).fetchall()
        finally:
            conn.close()
        image_model_names = [str(row["image_embedding_model"] or "").strip() for row in rows]
        configured_image_model = _normalize_image_embedding_model_name(_configured_image_embedding_model())
        if configured_image_model:
            image_model_names.append(configured_image_model)
        for model_name in sorted(set(image_model_names)):
            if not model_name:
                continue
            try:
                vector = _encode_text_for_image_embedding_model("K87 键盘实拍图", model_name)
                if not vector:
                    raise EmbeddingProviderError(f"empty image text embedding for {model_name}")
                models.append(_normalize_image_embedding_model_name(model_name))
            except Exception as exc:  # pragma: no cover - defensive startup warmup
                errors.append(f"{model_name}: {exc}")

        latency_ms = int((time.perf_counter() - started) * 1000)
        return {
            "status": "success" if not errors else "partial",
            "models": models,
            "errors": errors,
            "latency_ms": latency_ms,
        }

    def _record_retrieval(
        self,
        query: str,
        platform: str,
        shop_id: str,
        scene: str,
        result_count: int,
        latency_ms: int,
        mode: str,
    ) -> None:
        conn = self._connect()
        try:
            conn.execute(
                """
                INSERT INTO knowledge_retrieval_events
                (id, query, platform, shop_id, scene, result_count, latency_ms, mode, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (_new_id("retrieval"), query, platform, shop_id, scene, result_count, latency_ms, mode, _now_text()),
            )
            conn.commit()
        finally:
            conn.close()


def _row_to_dict(conn: sqlite3.Connection, row: sqlite3.Row | tuple[Any, ...]) -> dict[str, Any]:
    del conn
    if isinstance(row, sqlite3.Row):
        return dict(row)
    raise TypeError("knowledge store expected sqlite3.Row rows")


def _image_asset_search_text(
    title: str,
    asset_type: str,
    summary: str,
    ocr_text: str,
    tags: list[str],
    scenarios: list[str],
    risk_tags: list[str],
    suggested_reply: str,
) -> str:
    return "\n".join(
        part
        for part in [
            title,
            asset_type,
            summary,
            ocr_text,
            " ".join(tags),
            " ".join(scenarios),
            " ".join(risk_tags),
            suggested_reply,
        ]
        if str(part or "").strip()
    )


def _image_asset_row_to_dict(row: sqlite3.Row | None) -> dict[str, Any]:
    if row is None:
        return {}
    item = dict(row)
    item["tags"] = _json_loads_list(str(item.pop("tags_json", "[]") or "[]"))
    item["scenarios"] = _json_loads_list(str(item.pop("scenarios_json", "[]") or "[]"))
    item["risk_tags"] = _json_loads_list(str(item.pop("risk_tags_json", "[]") or "[]"))
    item["raw_analysis"] = _json_loads_object(str(item.pop("raw_analysis_json", "{}") or "{}"))
    item.pop("embedding", None)
    item.pop("image_embedding", None)
    return item


def _model_focus_terms(text: str) -> set[str]:
    normalized = text.upper()
    terms: set[str] = set()
    for match in re.findall(r"\bK\d{2,3}(?:\s*(?:PRO|MAX|LITE))?\b", normalized):
        terms.add(re.sub(r"\s+", " ", match).strip())
    return terms


def _model_base(term: str) -> str:
    match = re.search(r"K\d{2,3}", term.upper())
    return match.group(0) if match else term.upper().strip()


def _image_asset_combined_text(item: dict[str, Any]) -> str:
    return " ".join(
        str(part or "")
        for part in [
            item.get("title"),
            item.get("original_filename"),
            item.get("asset_type"),
            item.get("summary"),
            item.get("ocr_text"),
            " ".join(item.get("tags") or []),
            " ".join(item.get("scenarios") or []),
            item.get("suggested_reply"),
        ]
    )


def _explicit_model_mismatch(query: str, item: dict[str, Any]) -> bool:
    query_models = _model_focus_terms(query)
    if not query_models:
        return False
    asset_models = _model_focus_terms(_image_asset_combined_text(item))
    if not asset_models:
        return False
    query_bases = {_model_base(term) for term in query_models}
    asset_bases = {_model_base(term) for term in asset_models}
    return bool(query_bases and asset_bases and query_bases.isdisjoint(asset_bases))


def _image_rule_bonus(query: str, item: dict[str, Any]) -> float:
    query_terms = [term.lower() for term in _query_terms(query) if len(term.strip()) >= 2]
    if not query_terms:
        return 0.0
    title_text = f"{item.get('title') or ''} {item.get('original_filename') or ''}".lower()
    tag_text = " ".join(item.get("tags") or []).lower()
    scenario_text = " ".join(item.get("scenarios") or []).lower()
    body_text = f"{item.get('summary') or ''} {item.get('ocr_text') or ''} {item.get('suggested_reply') or ''}".lower()
    asset_type = str(item.get("asset_type") or "").lower()

    score = 0.0
    matched_terms: set[str] = set()
    for term in query_terms[:16]:
        if term in title_text:
            score += 0.18
            matched_terms.add(term)
        elif term in tag_text:
            score += 0.14
            matched_terms.add(term)
        elif term in body_text:
            score += 0.08
            matched_terms.add(term)
        elif term in scenario_text:
            score += 0.06
            matched_terms.add(term)

    query_models = _model_focus_terms(query)
    asset_models = _model_focus_terms(_image_asset_combined_text(item))
    if query_models and asset_models:
        query_bases = {_model_base(term) for term in query_models}
        asset_bases = {_model_base(term) for term in asset_models}
        if not query_bases.isdisjoint(asset_bases):
            score += 0.35

    if any(word in query for word in ("实拍", "图片", "照片", "图", "外观", "看下", "看看")):
        if any(word in asset_type for word in ("实拍", "商品", "规格", "尺寸", "安装")):
            score += 0.12
        if any(word in scenario_text for word in ("外观", "实拍", "尺寸", "安装", "规格")):
            score += 0.12

    if "图片名=" in body_text or "是否有图=是" in body_text:
        score += 0.08

    if matched_terms:
        score += min(0.15, 0.03 * len(matched_terms))
    return min(1.0, score)


def _rerank_image_candidates(query: str, candidates: list[dict[str, Any]]) -> dict[int, dict[str, Any]]:
    if not candidates:
        return {}
    if os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_RERANK_ENABLED", "1").strip() in {"0", "false", "False"}:
        return {}
    provider = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_RERANK_PROVIDER", "").strip().lower()
    if provider and provider not in {"ark", "doubao", "volcengine"}:
        return {}
    api_key = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_ANALYSIS_API_KEY", "").strip()
    model = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_ANALYSIS_MODEL", "").strip()
    base_url = os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_ANALYSIS_BASE_URL", "").strip() or ARK_IMAGE_ANALYSIS_DEFAULT_BASE_URL
    if not api_key or not model:
        return {}

    content: list[dict[str, Any]] = [
        {
            "type": "text",
            "text": (
                "请根据客户问题对候选客服图片素材重新排序，只输出合法 JSON。"
                "字段：{\"ranking\":[{\"index\":1,\"score\":0.0到1.0,\"reason\":\"简短原因\","
                "\"should_attach\":true或false,\"risk_notice\":\"风险提示或空\"}]}。"
                f"\n客户问题：{query}\n候选图片如下，index 从 1 开始："
            ),
        }
    ]
    for index, item in enumerate(candidates[:5], start=1):
        file_path = Path(str(item.get("file_path") or ""))
        if not file_path.is_file():
            continue
        info = {
            "index": index,
            "filename": item.get("original_filename") or "",
            "title": item.get("title") or "",
            "asset_type": item.get("asset_type") or "",
            "summary": item.get("summary") or "",
            "ocr_text": item.get("ocr_text") or "",
            "tags": item.get("tags") or [],
            "scenarios": item.get("scenarios") or [],
            "risk_tags": item.get("risk_tags") or [],
        }
        content.append({"type": "text", "text": json.dumps(info, ensure_ascii=False)})
        try:
            content.append({"type": "image_url", "image_url": {"url": _image_file_to_data_url(file_path)}})
        except Exception:
            continue
    if len(content) <= 1:
        return {}

    payload = {
        "model": model,
        "stream": False,
        "messages": [
            {
                "role": "system",
                "content": "你是电商客服图片素材匹配助手，只判断候选图片是否匹配客户当前问题，不要编造图片内容。",
            },
            {"role": "user", "content": content},
        ],
        "thinking": {"type": "disabled"},
        "temperature": 0.0,
    }
    request = Request(
        _chat_completions_url(base_url),
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        method="POST",
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"},
    )
    try:
        timeout = float(os.environ.get("AI_CUSTOMER_SERVICE_IMAGE_RERANK_TIMEOUT_SEC", "25") or "25")
        with urlopen(request, timeout=max(5.0, min(timeout, 60.0))) as response:
            body = response.read()
        content_text = _extract_chat_completion_content(body)
        parsed = _extract_json_object_from_text(content_text)
    except Exception as exc:
        logging.warning("Knowledge image rerank skipped query=%r error=%s", query, exc)
        return {}

    ranking = parsed.get("ranking")
    if not isinstance(ranking, list):
        return {}
    adjustments: dict[int, dict[str, Any]] = {}
    for rank, item in enumerate(ranking):
        if not isinstance(item, dict):
            continue
        try:
            index = int(item.get("index") or 0) - 1
        except (TypeError, ValueError):
            continue
        if index < 0 or index >= len(candidates):
            continue
        try:
            rerank_score = max(0.0, min(1.0, float(item.get("score") or 0.0)))
        except (TypeError, ValueError):
            rerank_score = 0.0
        score_boost = max(0.0, 0.25 * rerank_score - 0.03 * rank)
        adjustments[index] = {
            "score_boost": score_boost,
            "rerank_score": rerank_score,
            "reason": str(item.get("reason") or "").strip(),
            "match_type": "multimodal_rerank" if rerank_score >= 0.5 else "",
        }
    return adjustments


def _image_recommendation_reason(query: str, item: dict[str, Any]) -> str:
    scenarios = item.get("scenarios") or []
    if scenarios:
        return f"客户咨询与图片适用场景“{scenarios[0]}”匹配。"
    tags = item.get("tags") or []
    if tags:
        return f"客户咨询与图片标签“{tags[0]}”相关。"
    asset_type = str(item.get("asset_type") or "").strip()
    if asset_type and asset_type != "other":
        return f"客户咨询可能适合参考这张{asset_type}。"
    del query
    return "客户咨询与该图片素材的摘要内容相关。"


_store_lock = None
_store_instance: KnowledgeStore | None = None


def get_knowledge_store() -> KnowledgeStore:
    global _store_instance
    if _store_instance is None:
        _store_instance = KnowledgeStore()
    return _store_instance


def set_knowledge_store_for_tests(store: KnowledgeStore | None) -> None:
    global _store_instance
    _store_instance = store


